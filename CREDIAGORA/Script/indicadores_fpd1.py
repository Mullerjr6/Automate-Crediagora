import csv
import os
import re
import shutil
import threading
import time
import uuid
import zipfile
from contextlib import contextmanager
from copy import copy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

ABA_DADOS = "Tabela Conta"
COLUNAS_DADOS = 50
COLUNA_MODELO_INICIAL = 51
COLUNA_MODELO_FINAL = 90
COLUNA_FORMULA_LEGADA_FINAL = 91
COLUNAS_DATA = {
    "DataContabil", "DataEfetivacao", "Nascimento", "PrimeiroVencimento",
    "DataPagamentoPrimeiraParcela",
}
COLUNAS_TEXTO = {"Proposta", "Cpf", "Fone01", "Fone02"}
CABECALHOS_VAZIOS_PERMITIDOS = {23: "Fone01", 24: "Fone02"}
class AtualizacaoIndicadoresError(RuntimeError):
    def __init__(self, etapa, mensagem):
        self.etapa = etapa
        super().__init__(f"[INDICADORES][ERRO][{etapa}] {mensagem}")


@dataclass(frozen=True)
class ResultadoAtualizacaoIndicadores:
    csv_utilizado: Path
    xlsx_atualizado: Path
    total_linhas_csv: int
    registros_importados: int
    primeira_linha: int
    ultima_linha: int
    colunas_importadas: int
    colunas_modelo: int
    formulas_modelo: int


def _log(log, mensagem):
    log(f"[INDICADORES] {mensagem}")


@contextmanager
def _progresso_continuo(log, descricao, intervalo=20):
    terminou = threading.Event()
    inicio = time.monotonic()

    def acompanhar():
        while not terminou.wait(intervalo):
            _log(log, f"{descricao} em andamento... {time.monotonic() - inicio:.0f}s.")

    thread = threading.Thread(target=acompanhar, daemon=True)
    _log(log, f"{descricao} iniciado.")
    thread.start()
    try:
        yield
    finally:
        terminou.set()
        thread.join(timeout=1)
        _log(log, f"{descricao} finalizado em {time.monotonic() - inicio:.1f}s.")


def _validar_csv_estavel(caminho, tentativas=3, intervalo=1):
    caminho = Path(caminho)
    if not caminho.exists():
        raise AtualizacaoIndicadoresError("CSV", f"Arquivo nao encontrado: {caminho}")
    if caminho.suffix.casefold() != ".csv":
        raise AtualizacaoIndicadoresError("CSV", f"Extensao invalida: {caminho.suffix}")
    tamanhos = []
    for indice in range(tentativas):
        tamanhos.append(caminho.stat().st_size)
        if indice + 1 < tentativas:
            time.sleep(intervalo)
    if not tamanhos[-1]:
        raise AtualizacaoIndicadoresError("CSV", "O arquivo esta vazio.")
    if len(set(tamanhos)) != 1:
        raise AtualizacaoIndicadoresError("CSV", f"Download nao estabilizado: {tamanhos}")
    return caminho


def _ler_csv(caminho):
    amostra_bytes = caminho.read_bytes()[:65536]
    for codificacao in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            amostra = amostra_bytes.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise AtualizacaoIndicadoresError("CSV", "Codificacao nao identificada.")
    try:
        dialeto = csv.Sniffer().sniff(amostra, delimiters=";,\t|")
    except csv.Error as erro:
        raise AtualizacaoIndicadoresError("CSV", "Delimitador nao identificado.") from erro

    linhas = []
    with caminho.open("r", encoding=codificacao, newline="") as arquivo:
        for numero, linha in enumerate(csv.reader(arquivo, dialeto), start=1):
            if numero > 1 and not any(valor.strip() for valor in linha):
                continue
            if len(linha) < COLUNAS_DADOS:
                raise AtualizacaoIndicadoresError(
                    "CSV", f"Linha {numero}: {len(linha)} colunas; esperado: 50."
                )
            if any(valor.strip() for valor in linha[COLUNAS_DADOS:]):
                raise AtualizacaoIndicadoresError(
                    "CSV", f"Linha {numero} possui dados depois da coluna AX."
                )
            linhas.append(linha[:COLUNAS_DADOS])
    if len(linhas) < 2:
        raise AtualizacaoIndicadoresError("CSV", "CSV sem registros de dados.")
    return linhas[0], linhas[1:], codificacao, dialeto.delimiter


def _converter(valor, cabecalho):
    if valor == "":
        return None
    if cabecalho in COLUNAS_TEXTO:
        return valor
    if cabecalho in COLUNAS_DATA:
        try:
            return datetime.strptime(valor, "%d/%m/%Y")
        except ValueError as erro:
            raise AtualizacaoIndicadoresError(
                "Conversao", f"Data invalida em {cabecalho}: {valor!r}"
            ) from erro
    numero = valor.strip().replace(".", "").replace(",", ".")
    if re.fullmatch(r"[-+]?\d+", numero):
        return int(numero)
    if re.fullmatch(r"[-+]?(?:\d+\.\d*|\d*\.\d+)", numero):
        return float(numero)
    return valor


def _validar_cabecalhos(cabecalho_csv, cabecalhos_xlsx):
    for posicao, (csv_nome, xlsx_nome) in enumerate(
        zip(cabecalho_csv, cabecalhos_xlsx), start=1
    ):
        if csv_nome == xlsx_nome:
            continue
        if (
            not csv_nome
            and CABECALHOS_VAZIOS_PERMITIDOS.get(posicao) == xlsx_nome
        ):
            continue
        raise AtualizacaoIndicadoresError(
            "CSV",
            f"Cabecalho divergente na coluna {posicao}: "
            f"CSV={csv_nome!r}; XLSX={xlsx_nome!r}.",
        )
    return cabecalhos_xlsx


def _traduzir(modelo, origem, destino):
    """Traduz uma formula como o preenchimento da alca do Excel.

    As formulas do modelo possuem referencias externas, intervalos e referencias
    absolutas. Uma troca textual de numeros de linha nao cobre todos esses casos.
    """
    if isinstance(modelo, str) and modelo.startswith("="):
        try:
            return Translator(modelo, origin=origem).translate_formula(destino)
        except Exception as erro:
            raise AtualizacaoIndicadoresError(
                "Formulas", f"Nao foi possivel traduzir {origem} para {destino}."
            ) from erro
    return modelo


def reparar_formulas_indicadores_fpd1(xlsx_destino, pasta_backup, log=print):
    """Reconstroi apenas AY:CL de uma planilha ja atualizada.

    Os dados de A:AX permanecem intactos. A linha 2 e o modelo e cada linha
    seguinte recebe a traducao equivalente ao preenchimento de formula no Excel.
    """
    inicio = time.monotonic()
    xlsx_destino = Path(xlsx_destino)
    pasta_backup = Path(pasta_backup)
    if not xlsx_destino.exists():
        raise AtualizacaoIndicadoresError("Reparo", f"Destino nao encontrado: {xlsx_destino}")

    wb = None
    temporario = xlsx_destino.with_name(f".{xlsx_destino.stem}.{uuid.uuid4().hex}.tmp.xlsx")
    backup = None
    substituido = False
    try:
        with _progresso_continuo(log, "Carregamento do XLSX para reparo"):
            wb = load_workbook(xlsx_destino, data_only=False, keep_links=False)
        if ABA_DADOS not in wb.sheetnames:
            raise AtualizacaoIndicadoresError("Reparo", f"Aba {ABA_DADOS!r} ausente.")
        ws = wb[ABA_DADOS]
        ultima = ws.max_row
        if ultima < 2:
            raise AtualizacaoIndicadoresError("Reparo", "Planilha sem linha-modelo para formulas.")

        modelo = tuple(
            ws.cell(2, coluna).value
            for coluna in range(COLUNA_MODELO_INICIAL, COLUNA_MODELO_FINAL + 1)
        )
        formulas = sum(isinstance(valor, str) and valor.startswith("=") for valor in modelo)
        if not formulas:
            raise AtualizacaoIndicadoresError("Reparo", "AY2:CL2 sem formulas-modelo.")

        _log(log, f"Recompondo AY:CL da linha 2 ate a linha {ultima}.")
        with _progresso_continuo(log, "Reconstrucao das formulas"):
            for linha in range(3, ultima + 1):
                for coluna, valor in enumerate(modelo, start=COLUNA_MODELO_INICIAL):
                    destino = f"{get_column_letter(coluna)}{linha}"
                    celula = ws.cell(linha, coluna)
                    if not isinstance(celula, Cell):
                        raise AtualizacaoIndicadoresError(
                            "Reparo", f"Celula mesclada inesperada em {destino}."
                        )
                    celula.value = _traduzir(
                        valor, f"{get_column_letter(coluna)}2", destino
                    )
                if linha % 1000 == 0:
                    _log(log, f"Reparo de formulas: linha {linha} de {ultima}.")

        # CM era controlada pela automacao antiga. Remove somente formulas
        # residuais para nao apagar um eventual valor fixo preenchido pelo usuario.
        for linha in range(2, ultima + 1):
            celula_legada = ws.cell(linha, COLUNA_FORMULA_LEGADA_FINAL)
            if not isinstance(celula_legada, Cell):
                raise AtualizacaoIndicadoresError(
                    "Reparo", f"Celula mesclada inesperada em {celula_legada.coordinate}."
                )
            if isinstance(celula_legada.value, str) and celula_legada.value.startswith("="):
                celula_legada.value = None

        linhas_amostra = sorted({2, min(3, ultima), (ultima + 2) // 2, ultima})
        for linha in linhas_amostra:
            for coluna, valor in enumerate(modelo, start=COLUNA_MODELO_INICIAL):
                origem = f"{get_column_letter(coluna)}2"
                destino = f"{get_column_letter(coluna)}{linha}"
                esperado = _traduzir(valor, origem, destino)
                if ws.cell(linha, coluna).value != esperado:
                    raise AtualizacaoIndicadoresError(
                        "Reparo", f"Formula nao confirmada em {destino}."
                    )

        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
        with _progresso_continuo(log, "Salvamento do XLSX reparado"):
            links_externos = getattr(wb, "_external_links", None)
            if isinstance(links_externos, list):
                links_externos.clear()
            wb.save(temporario)
        wb.close()
        wb = None
        _validar_arquivo_salvo(temporario)

        pasta_backup.mkdir(parents=True, exist_ok=True)
        backup = pasta_backup / (
            f"{xlsx_destino.stem}_antes_reparo_formulas_"
            f"{datetime.now().strftime('%Y%m%d_%H%M%S')}{xlsx_destino.suffix}"
        )
        shutil.copy2(xlsx_destino, backup)
        os.replace(temporario, xlsx_destino)
        substituido = True
        _validar_arquivo_salvo(xlsx_destino)
        substituido = False
        _log(log, f"Backup preservado: {backup}")
        _log(log, f"Formulas reparadas: AY2:CL{ultima} ({formulas} formulas-modelo).")
        _log(log, f"Duracao do reparo: {time.monotonic() - inicio:.1f}s.")
        return xlsx_destino, backup, ultima, formulas
    except AtualizacaoIndicadoresError:
        raise
    except PermissionError as erro:
        raise AtualizacaoIndicadoresError(
            "Reparo", "Planilha em uso no Excel. Feche a planilha TESTE e execute novamente."
        ) from erro
    except Exception as erro:
        raise AtualizacaoIndicadoresError("Reparo", str(erro)) from erro
    finally:
        if wb is not None:
            wb.close()
        temporario.unlink(missing_ok=True)
        if substituido and backup is not None and backup.exists():
            os.replace(backup, xlsx_destino)


def _validar_planilha(ws, cabecalhos, modelo, registros, ultima, antiga_ultima):
    atuais = tuple(ws.cell(1, c).value for c in range(1, COLUNAS_DADOS + 1))
    if atuais != cabecalhos:
        raise AtualizacaoIndicadoresError("Validacao", "Cabecalhos foram alterados.")
    if ws.cell(2, 1).value != registros[0][0] or ws.cell(2, 50).value != registros[0][-1]:
        raise AtualizacaoIndicadoresError("Validacao", "Primeiro registro nao confere.")
    if ws.cell(ultima, 1).value != registros[-1][0]:
        raise AtualizacaoIndicadoresError("Validacao", "Ultimo registro nao confere.")
    for coluna, valor in enumerate(modelo, start=COLUNA_MODELO_INICIAL):
        origem = f"{get_column_letter(coluna)}2"
        if ws.cell(2, coluna).value != valor:
            raise AtualizacaoIndicadoresError("Validacao", f"Modelo alterado em {origem}.")
        destino = ws.cell(ultima, coluna)
        coordenada_destino = f"{get_column_letter(coluna)}{ultima}"
        if destino.value != _traduzir(valor, origem, coordenada_destino):
            raise AtualizacaoIndicadoresError("Validacao", f"Modelo incorreto em {coordenada_destino}.")
    for linha in range(ultima + 1, antiga_ultima + 1):
        for coluna in range(1, COLUNA_MODELO_FINAL + 1):
            if ws.cell(linha, coluna).value is not None:
                raise AtualizacaoIndicadoresError(
                    "Validacao", f"Conteudo antigo em {ws.cell(linha, coluna).coordinate}."
                )


def _validar_arquivo_salvo(caminho):
    """Confere se o XLSX salvo e um ZIP completo e legivel sem reabrir planilhas."""
    if not caminho.exists() or caminho.stat().st_size <= 0:
        raise AtualizacaoIndicadoresError("Pos-salvamento", "XLSX ausente ou vazio.")
    try:
        with zipfile.ZipFile(caminho) as arquivo:
            obrigatorios = {"[Content_Types].xml", "xl/workbook.xml"}
            if not obrigatorios.issubset(arquivo.namelist()) or arquivo.testzip() is not None:
                raise AtualizacaoIndicadoresError("Pos-salvamento", "XLSX incompleto ou corrompido.")
    except zipfile.BadZipFile as erro:
        raise AtualizacaoIndicadoresError("Pos-salvamento", "XLSX corrompido.") from erro


def atualizar_indicadores_fpd1(csv_atual, xlsx_destino, log=print):
    inicio = time.monotonic()
    csv_atual = _validar_csv_estavel(csv_atual)
    xlsx_destino = Path(xlsx_destino)
    if not xlsx_destino.exists():
        raise AtualizacaoIndicadoresError("XLSX", f"Destino nao encontrado: {xlsx_destino}")
    cabecalho_csv, linhas_csv, codificacao, delimitador = _ler_csv(csv_atual)
    _log(log, f"CSV utilizado: {csv_atual}")
    _log(log, f"Codificacao={codificacao}; delimitador={delimitador!r}.")
    _log(log, f"Total de linhas do CSV: {len(linhas_csv) + 1}.")

    wb = None
    temporario = xlsx_destino.with_name(f".{xlsx_destino.stem}.{uuid.uuid4().hex}.tmp.xlsx")
    backup = xlsx_destino.with_name(f".{xlsx_destino.stem}.{uuid.uuid4().hex}.bak.xlsx")
    precisa_restaurar = False
    try:
        with _progresso_continuo(log, "Carregamento do XLSX"):
            wb = load_workbook(xlsx_destino, data_only=False, keep_links=False)
        if ABA_DADOS not in wb.sheetnames:
            raise AtualizacaoIndicadoresError("XLSX", f"Aba {ABA_DADOS!r} ausente.")
        ws = wb[ABA_DADOS]
        antiga_ultima = ws.max_row
        cabecalhos = tuple(ws.cell(1, c).value for c in range(1, 51))
        cabecalhos_efetivos = _validar_cabecalhos(cabecalho_csv, cabecalhos)
        with _progresso_continuo(log, "Conferencia da quantidade de registros"):
            registros_anteriores = sum(
                any(valor is not None and str(valor).strip() for valor in linha)
                for linha in ws.iter_rows(
                    min_row=2, max_row=antiga_ultima,
                    min_col=1, max_col=COLUNAS_DADOS, values_only=True,
                )
            )
        _log(log, f"Comparacao de registros: CSV={len(linhas_csv)}; XLSX={registros_anteriores}.")
        if len(linhas_csv) < registros_anteriores:
            mensagem = (
                f"CSV possui {len(linhas_csv)} registros, menos que os "
                f"{registros_anteriores} registros existentes no XLSX "
                f"(diferenca: {registros_anteriores - len(linhas_csv)}). "
                "Exportacao possivelmente incompleta. Atualizacao bloqueada; "
                "XLSX preservado e CSV mantido para diagnostico."
            )
            _log(log, mensagem)
            raise AtualizacaoIndicadoresError("Quantidade de registros", mensagem)
        modelo = tuple(
            ws.cell(2, coluna).value
            for coluna in range(COLUNA_MODELO_INICIAL, COLUNA_MODELO_FINAL + 1)
        )
        formulas = sum(isinstance(v, str) and v.startswith("=") for v in modelo)
        if not formulas:
            raise AtualizacaoIndicadoresError("Formulas", "AY2:CL2 sem formulas-modelo.")
        registros = [
            tuple(_converter(v, cabecalhos_efetivos[i]) for i, v in enumerate(linha))
            for linha in linhas_csv
        ]
        ultima = len(registros) + 1
        limite = max(antiga_ultima, ultima)
        estilos_modelo = [
            copy(ws.cell(2, coluna)._style)
            for coluna in range(1, COLUNA_MODELO_FINAL + 1)
        ]

        _log(log, f"Removendo base anterior ate a linha {limite}.")
        coordenadas_antigas = [
            coordenada for coordenada in ws._cells
            if (
                coordenada[0] >= 2
                and coordenada[1] <= COLUNA_FORMULA_LEGADA_FINAL
            )
        ]
        for coordenada in coordenadas_antigas:
            del ws._cells[coordenada]
        _log(log, "Limpeza da base anterior concluida.")

        _log(log, f"Gravando {len(registros)} registros e formulas.")
        for numero_linha, registro in enumerate(registros, start=2):
            for coluna, valor in enumerate(registro, start=1):
                celula = ws.cell(numero_linha, coluna, valor)
                if numero_linha > antiga_ultima:
                    celula._style = copy(estilos_modelo[coluna - 1])
                if isinstance(valor, datetime):
                    celula.number_format = "dd/mm/yyyy"
            for coluna, valor in enumerate(modelo, start=COLUNA_MODELO_INICIAL):
                celula = ws.cell(numero_linha, coluna)
                if not isinstance(celula, Cell):
                    raise AtualizacaoIndicadoresError(
                        "Formulas", f"Celula mesclada inesperada em {celula.coordinate}."
                    )
                celula.value = _traduzir(valor, f"{get_column_letter(coluna)}2", celula.coordinate)
                if numero_linha > antiga_ultima:
                    celula._style = copy(estilos_modelo[coluna - 1])
            if numero_linha % 1000 == 0:
                _log(log, f"Gravacao: linha {numero_linha} de {ultima}.")
        _log(log, "Dados e formulas preenchidos em memoria.")

        _validar_planilha(ws, cabecalhos, modelo, registros, ultima, antiga_ultima)
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
        with _progresso_continuo(log, "Salvamento do XLSX temporario"):
            links_externos = getattr(wb, "_external_links", None)
            if isinstance(links_externos, list):
                links_externos.clear()
            wb.save(temporario)
        wb.close()
        wb = None
        with _progresso_continuo(log, "Conferencia do XLSX temporario"):
            _validar_arquivo_salvo(temporario)
        shutil.copy2(xlsx_destino, backup)
        os.replace(temporario, xlsx_destino)
        precisa_restaurar = True
        with _progresso_continuo(log, "Conferencia final do XLSX"):
            _validar_arquivo_salvo(xlsx_destino)
        precisa_restaurar = False
        backup.unlink(missing_ok=True)

        resultado = ResultadoAtualizacaoIndicadores(
            csv_atual, xlsx_destino, len(linhas_csv) + 1, len(registros), 2,
            ultima, 50, len(modelo), formulas,
        )
        _log(log, f"Registros importados: {len(registros)}.")
        _log(log, f"Ultima linha escrita: {ultima}.")
        _log(log, "Colunas importadas: 50 (A:AX).")
        _log(log, f"Colunas-modelo: {len(modelo)} (AY:CL); formulas-modelo: {formulas}.")
        primeira_formula = next(v for v in modelo if isinstance(v, str) and v.startswith("="))
        _log(log, f"Primeira formula: {primeira_formula}")
        _log(log, f"Ultima linha com formulas: {ultima}.")
        _log(log, f"Duracao da atualizacao: {time.monotonic() - inicio:.1f}s.")
        _log(log, "ATUALIZACAO INDICADORES CONCLUIDA COM SUCESSO")
        return resultado
    except AtualizacaoIndicadoresError:
        raise
    except Exception as erro:
        raise AtualizacaoIndicadoresError("Atualizacao", str(erro)) from erro
    finally:
        if wb is not None:
            wb.close()
        temporario.unlink(missing_ok=True)
        if precisa_restaurar and backup.exists():
            os.replace(backup, xlsx_destino)
        else:
            backup.unlink(missing_ok=True)
