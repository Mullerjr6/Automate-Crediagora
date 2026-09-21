import ctypes
import json
import os
import re
import shutil
import subprocess
import tempfile
import time
import unicodedata
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

from ctypes import wintypes
from openpyxl import load_workbook


def _configurar_dpi_windows():
    """Mantem Win32, capturas e mouse no mesmo sistema de coordenadas."""
    try:
        # PROCESS_PER_MONITOR_DPI_AWARE. Precisa ocorrer antes de PIL/pywinauto.
        ctypes.windll.shcore.SetProcessDpiAwareness(2)
    except Exception:
        try:
            ctypes.windll.user32.SetProcessDPIAware()
        except Exception:
            pass


_configurar_dpi_windows()

from PIL import Image, ImageChops, ImageGrab, ImageOps


LOJAS = ("TUC", "GMI", "DMC", "RGO", "CMI", "CAT", "SPM", "SCR")
LINHAS_EXCEL = dict(zip(LOJAS, range(19, 27)))


class GestorError(RuntimeError):
    def __init__(self, etapa, mensagem, hwnd=None):
        self.etapa = etapa
        self.hwnd = hwnd
        super().__init__(f"[GESTOR][ERRO][{etapa}] {mensagem}")


@dataclass(frozen=True)
class GestorConfig:
    remoto_usuario: str
    usuario: str
    senha: str
    arquivo_excel: Path
    pasta_logs: Path
    pasta_erros: Path
    pasta_backup: Path
    atalho: Path = Path.home() / "Desktop" / "GESTOR NUVEM.lnk"
    timeout_janela: int = 60
    timeout_relatorio: int = 120
    intervalo_tecla: float = 0.12

    def validar(self):
        ausentes = []
        if not str(self.remoto_usuario).strip():
            ausentes.append("GESTOR_REMOTO_USUARIO")
        if not str(self.usuario).strip():
            ausentes.append("GESTOR_USUARIO")
        if not str(self.senha).strip():
            ausentes.append("GESTOR_SENHA")
        if ausentes:
            raise EnvironmentError(
                "Credenciais Gestor nao configuradas: " + ", ".join(ausentes)
            )
        if not self.atalho.exists():
            raise FileNotFoundError(f"Atalho do Gestor nao encontrado: {self.atalho}")
        if not self.arquivo_excel.exists():
            raise FileNotFoundError(
                f"Planilha de metas do Gestor nao encontrada: {self.arquivo_excel}"
            )


@dataclass(frozen=True)
class ResultadoGestor:
    data_relatorio: str
    valores: dict
    total_relatorio: Decimal
    soma_lojas: Decimal
    arquivo_excel: Path
    screenshot: Path
    paginas: int
    duracao_segundos: float


def normalizar(texto):
    texto = unicodedata.normalize("NFD", str(texto or ""))
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", texto.replace("\ufffd", "")).casefold().strip()


def log_gestor(log, mensagem):
    log(f"[GESTOR] {mensagem}")


def numero_brasileiro(valor):
    texto = str(valor).strip().replace("R$", "").replace(" ", "")
    if not re.fullmatch(r"-?\d{1,3}(?:\.\d{3})*,\d{2}", texto):
        raise ValueError(f"Numero brasileiro invalido: {valor!r}")
    try:
        return Decimal(texto.replace(".", "").replace(",", "."))
    except InvalidOperation as erro:
        raise ValueError(f"Numero brasileiro invalido: {valor!r}") from erro


def formatar_brasileiro(valor):
    valor = Decimal(str(valor)).quantize(Decimal("0.01"))
    inteiro, centavos = f"{valor:.2f}".split(".")
    grupos = []
    while inteiro:
        grupos.append(inteiro[-3:])
        inteiro = inteiro[:-3]
    return ".".join(reversed(grupos)) + "," + centavos


def validar_total_relatorio(valores, total_relatorio):
    faltantes = [loja for loja in LOJAS if loja not in valores]
    extras = [loja for loja in valores if loja not in LOJAS]
    if faltantes or extras or len(valores) != len(LOJAS):
        raise GestorError(
            "Validacao do relatorio",
            f"Lojas invalidas. Faltantes={faltantes}; extras={extras}.",
        )
    soma = sum((Decimal(str(valores[loja])) for loja in LOJAS), Decimal("0"))
    total = Decimal(str(total_relatorio))
    if soma.quantize(Decimal("0.01")) != total.quantize(Decimal("0.01")):
        raise GestorError(
            "Validacao do relatorio",
            "Soma das lojas nao confere com o TOTAL GERAL: "
            f"{formatar_brasileiro(soma)} != {formatar_brasileiro(total)}.",
        )
    return soma.quantize(Decimal("0.01"))


class RECT(ctypes.Structure):
    _fields_ = [
        ("left", wintypes.LONG),
        ("top", wintypes.LONG),
        ("right", wintypes.LONG),
        ("bottom", wintypes.LONG),
    ]


class POINT(ctypes.Structure):
    _fields_ = [("x", wintypes.LONG), ("y", wintypes.LONG)]


class MONITORINFOEXW(ctypes.Structure):
    _fields_ = [
        ("cbSize", wintypes.DWORD),
        ("rcMonitor", RECT),
        ("rcWork", RECT),
        ("dwFlags", wintypes.DWORD),
        ("szDevice", wintypes.WCHAR * 32),
    ]


def _rect_tupla(rect):
    return rect.left, rect.top, rect.right, rect.bottom


def _monitor_principal():
    user32 = ctypes.windll.user32
    hmonitor = user32.MonitorFromPoint(POINT(0, 0), 1)
    info = MONITORINFOEXW()
    info.cbSize = ctypes.sizeof(info)
    if not user32.GetMonitorInfoW(hmonitor, ctypes.byref(info)):
        raise GestorError("Janela", "Nao foi possivel consultar o monitor principal.")
    return {
        "handle": int(hmonitor),
        "dispositivo": info.szDevice,
        "retangulo": _rect_tupla(info.rcMonitor),
        "area_trabalho": _rect_tupla(info.rcWork),
    }


def _retangulo_visivel(hwnd):
    rect = RECT()
    if not ctypes.windll.user32.GetWindowRect(hwnd, ctypes.byref(rect)):
        raise GestorError("Janela", f"GetWindowRect falhou para hwnd={hwnd}.")
    try:
        visivel = RECT()
        resultado = ctypes.windll.dwmapi.DwmGetWindowAttribute(
            hwnd, 9, ctypes.byref(visivel), ctypes.sizeof(visivel)
        )
        if resultado == 0:
            return _rect_tupla(visivel)
    except Exception:
        pass
    return _rect_tupla(rect)


def _enumerar_janelas():
    user32 = ctypes.windll.user32
    janelas = []
    callback_type = ctypes.WINFUNCTYPE(wintypes.BOOL, wintypes.HWND, wintypes.LPARAM)

    @callback_type
    def callback(hwnd, _dados):
        if not user32.IsWindowVisible(hwnd):
            return True
        tamanho = user32.GetWindowTextLengthW(hwnd)
        if tamanho <= 0:
            return True
        titulo = ctypes.create_unicode_buffer(tamanho + 1)
        classe = ctypes.create_unicode_buffer(256)
        user32.GetWindowTextW(hwnd, titulo, len(titulo))
        user32.GetClassNameW(hwnd, classe, len(classe))
        try:
            retangulo = _retangulo_visivel(int(hwnd))
        except Exception:
            return True
        janelas.append(
            {
                "hwnd": int(hwnd),
                "titulo": titulo.value,
                "classe": classe.value,
                "retangulo": retangulo,
            }
        )
        return True

    user32.EnumWindows(callback, 0)
    return janelas


def _handles_janelas():
    return {janela["hwnd"] for janela in _enumerar_janelas()}


def _estado_janelas():
    return {
        janela["hwnd"]: normalizar(janela["titulo"])
        for janela in _enumerar_janelas()
    }


def _janela_ativa():
    return int(ctypes.windll.user32.GetForegroundWindow() or 0)


def aguardar_janela(
    termos,
    timeout,
    novos_desde=None,
    classe=None,
    etapa="Janela",
):
    if isinstance(termos, str):
        termos = (termos,)
    termos = tuple(normalizar(termo) for termo in termos)
    classe_normalizada = normalizar(classe) if classe else None
    fim = time.monotonic() + timeout
    ultimo = []
    while time.monotonic() < fim:
        candidatos = []
        for janela in _enumerar_janelas():
            titulo = normalizar(janela["titulo"])
            nome_classe = normalizar(janela["classe"])
            if not all(termo in titulo for termo in termos):
                continue
            if classe_normalizada and classe_normalizada not in nome_classe:
                continue
            candidatos.append(janela)
        if novos_desde is not None:
            handles_anteriores = (
                set(novos_desde)
                if isinstance(novos_desde, dict)
                else set(novos_desde)
            )
            novos = [j for j in candidatos if j["hwnd"] not in handles_anteriores]
            if novos:
                return novos[0]
            if isinstance(novos_desde, dict):
                alterados = [
                    j
                    for j in candidatos
                    if novos_desde.get(j["hwnd"]) != normalizar(j["titulo"])
                ]
                if alterados:
                    ativa = _janela_ativa()
                    alterados.sort(key=lambda j: j["hwnd"] != ativa)
                    return alterados[0]
        elif candidatos:
            return candidatos[0]
        ultimo = [
            f"{j['titulo']} [hwnd={j['hwnd']}; classe={j['classe']}; rect={j['retangulo']}]"
            for j in _enumerar_janelas()
        ]
        time.sleep(0.5)
    raise GestorError(
        etapa,
        f"Janela nao encontrada: {' + '.join(termos)}. Visiveis={ultimo[:12]}",
        hwnd=_janela_ativa() or None,
    )


def _janela_existe(hwnd):
    return bool(ctypes.windll.user32.IsWindow(hwnd)) and bool(
        ctypes.windll.user32.IsWindowVisible(hwnd)
    )


def _aguardar_fechar(hwnd, timeout=20):
    fim = time.monotonic() + timeout
    while time.monotonic() < fim:
        if not _janela_existe(hwnd):
            return
        time.sleep(0.5)
    raise GestorError(
        "Janela", f"A janela hwnd={hwnd} nao fechou no tempo esperado.", hwnd=hwnd
    )


def _fechar_janela_gestor(hwnd, timeout=20):
    if not hwnd or not _janela_existe(hwnd):
        return
    user32 = ctypes.windll.user32
    user32.PostMessageW(hwnd, 0x0010, 0, 0)
    try:
        _aguardar_fechar(hwnd, timeout=timeout)
    except GestorError as erro:
        raise GestorError(
            "Janela",
            f"Nao foi possivel fechar a janela do Gestor de Vendas: {erro}",
            hwnd=hwnd,
        ) from erro


def _retangulo_contido(retangulo, limite, tolerancia=20):
    return (
        retangulo[0] >= limite[0] - tolerancia
        and retangulo[1] >= limite[1] - tolerancia
        and retangulo[2] <= limite[2] + tolerancia
        and retangulo[3] <= limite[3] + tolerancia
    )


def preparar_janela_gestor(hwnd, log, mover=True):
    user32 = ctypes.windll.user32
    monitor = _monitor_principal()
    retangulo = _retangulo_visivel(hwnd)
    antes = {
        "retangulo": retangulo,
        "maximizada": bool(user32.IsZoomed(hwnd)),
    }
    contida = _retangulo_contido(retangulo, monitor["retangulo"])
    if not contida and mover:
        # O Gestor pode permanecer no tamanho padrao. Apenas restaure e mova
        # quando alguma parte estiver fora do monitor principal.
        user32.ShowWindow(hwnd, 9)  # SW_RESTORE
        time.sleep(0.3)
        atual = _retangulo_visivel(hwnd)
        wl, wt, wr, wb = monitor["area_trabalho"]
        largura_disponivel = wr - wl
        altura_disponivel = wb - wt
        largura = min(max(atual[2] - atual[0], 320), largura_disponivel)
        altura = min(max(atual[3] - atual[1], 200), altura_disponivel)
        x = wl + max(0, (largura_disponivel - largura) // 2)
        y = wt + max(0, (altura_disponivel - altura) // 2)
        flags = 0x0004 | 0x0010 | 0x0040  # NOZORDER | NOACTIVATE | SHOWWINDOW
        if not user32.SetWindowPos(hwnd, 0, x, y, largura, altura, flags):
            raise GestorError(
                "Janela",
                f"Nao foi possivel mover hwnd={hwnd} para o monitor principal.",
                hwnd=hwnd,
            )
        time.sleep(0.5)
        retangulo = _retangulo_visivel(hwnd)
        contida = _retangulo_contido(retangulo, monitor["retangulo"])
    depois = {
        "retangulo": retangulo,
        "maximizada": bool(user32.IsZoomed(hwnd)),
        "monitor": monitor["dispositivo"],
        "contida": contida,
    }
    if not contida:
        raise GestorError(
            "Janela",
            f"Janela nao ficou contida no monitor principal: {depois}",
            hwnd=hwnd,
        )
    log_gestor(log, f"Janela padrao validada no {monitor['dispositivo']}: {depois}")
    return {"antes": antes, "depois": depois}


def _ativar(hwnd):
    user32 = ctypes.windll.user32
    if _janela_ativa() == hwnd:
        return
    classe = ctypes.create_unicode_buffer(256)
    user32.GetClassNameW(hwnd, classe, len(classe))
    if classe.value == "RAIL_WINDOW":
        # Nao maximizar nem alterar o tamanho das RemoteApps. O acoplamento
        # temporario das threads permite recuperar foco quando o terminal foi
        # a ultima janela usada, sem enviar teclas ou cliques auxiliares.
        foreground = _janela_ativa()
        thread_atual = ctypes.windll.kernel32.GetCurrentThreadId()
        thread_foreground = user32.GetWindowThreadProcessId(foreground, None)
        anexou = False
        try:
            if thread_foreground and thread_foreground != thread_atual:
                anexou = bool(user32.AttachThreadInput(thread_atual, thread_foreground, True))
            user32.BringWindowToTop(hwnd)
            user32.SetForegroundWindow(hwnd)
        finally:
            if anexou:
                user32.AttachThreadInput(thread_atual, thread_foreground, False)
        if _janela_ativa() != hwnd:
            raise GestorError(
                "Janela",
                f"Nao foi possivel trazer a janela remota para frente: hwnd={hwnd}.",
                hwnd=hwnd,
            )
        time.sleep(0.25)
        return
    user32.ShowWindow(hwnd, 5)  # SW_SHOW, preserva o estado atual
    SWP_NOSIZE = 0x0001
    SWP_NOMOVE = 0x0002
    SWP_SHOWWINDOW = 0x0040
    flags = SWP_NOSIZE | SWP_NOMOVE | SWP_SHOWWINDOW
    user32.SetWindowPos(hwnd, -1, 0, 0, 0, 0, flags)
    user32.BringWindowToTop(hwnd)
    user32.SetForegroundWindow(hwnd)
    user32.SetWindowPos(hwnd, -2, 0, 0, 0, 0, flags)
    if user32.GetForegroundWindow() != hwnd:
        try:
            user32.SwitchToThisWindow(hwnd, True)
        except Exception:
            pass
    time.sleep(0.25)


def _ponto_relativo(hwnd, x, y):
    left, top, right, bottom = _retangulo_visivel(hwnd)
    return left + int((right - left) * x), top + int((bottom - top) * y)


def _clicar(hwnd, x, y, cliques=1, intervalo=0.15):
    import pyautogui

    _ativar(hwnd)
    px, py = _ponto_relativo(hwnd, x, y)
    pyautogui.click(px, py, clicks=cliques, interval=intervalo)


def _mover(hwnd, x, y):
    import pyautogui

    _ativar(hwnd)
    px, py = _ponto_relativo(hwnd, x, y)
    pyautogui.moveTo(px, py, duration=0.25)


def _preencher_por_coordenada(hwnd, x, y, texto, intervalo=0.12, duplo=True):
    import pyautogui

    _clicar(hwnd, x, y, cliques=2 if duplo else 1)
    time.sleep(0.8)
    pyautogui.hotkey("ctrl", "a")
    time.sleep(0.25)
    pyautogui.press("backspace")
    time.sleep(0.25)
    pyautogui.write(str(texto), interval=intervalo)


def _janela_pywinauto(hwnd, backend="uia"):
    from pywinauto import Application

    return Application(backend=backend).connect(handle=hwnd).window(handle=hwnd)


def _controle_texto(janela, texto, tipos=("Button", "Text", "ListItem", "MenuItem")):
    alvo = normalizar(texto)
    parciais = []
    for tipo in tipos:
        try:
            controles = janela.descendants(control_type=tipo)
        except Exception:
            continue
        for controle in controles:
            try:
                atual = normalizar(controle.window_text())
                if atual == alvo:
                    return controle
                if alvo and alvo in atual:
                    parciais.append(controle)
            except Exception:
                continue
    if parciais:
        return parciais[0]
    raise GestorError("Interface", f"Controle nao encontrado: {texto}")


def _invocar_controle(controle):
    try:
        controle.invoke()
    except Exception:
        controle.click_input()


def conectar_remoteapp(config, log):
    log_gestor(log, "Abrindo o atalho GESTOR NUVEM...")
    anteriores = _estado_janelas()
    os.startfile(str(config.atalho))
    acesso = aguardar_janela(
        "acesso remoto",
        config.timeout_janela,
        novos_desde=anteriores,
        etapa="Acesso Remoto",
    )
    janela = _janela_pywinauto(acesso["hwnd"])
    edits = []
    try:
        edits = janela.descendants(control_type="Edit")
    except Exception:
        pass
    if edits:
        usuario = edits[0]
        try:
            atual = usuario.get_value()
        except Exception:
            atual = usuario.window_text()
        if normalizar(atual) != normalizar(config.remoto_usuario):
            usuario.set_focus()
            usuario.type_keys("^a{BACKSPACE}" + config.remoto_usuario, set_foreground=True)
    botao = _controle_texto(janela, "Conectar", tipos=("Button",))
    anteriores = _estado_janelas()
    _invocar_controle(botao)
    try:
        seguranca = aguardar_janela(
            "remoteapp",
            8,
            novos_desde=anteriores,
            classe="TscShellContainerClass",
            etapa="Aviso de seguranca RemoteApp",
        )
    except GestorError:
        # O launcher AutoIt pode ignorar o primeiro clique enquanto restaura
        # a senha salva. Reative a janela e confirme novamente uma unica vez.
        _ativar(acesso["hwnd"])
        _invocar_controle(botao)
        seguranca = aguardar_janela(
            "remoteapp",
            config.timeout_janela,
            novos_desde=anteriores,
            classe="TscShellContainerClass",
            etapa="Aviso de seguranca RemoteApp",
        )
    janela_seg = _janela_pywinauto(seguranca["hwnd"])
    fim_checkboxes = time.monotonic() + min(20, config.timeout_janela)
    checkboxes = []
    while time.monotonic() < fim_checkboxes:
        try:
            checkboxes = janela_seg.descendants(control_type="CheckBox")
        except Exception:
            checkboxes = []
        if any(caixa.is_visible() for caixa in checkboxes):
            break
        time.sleep(0.5)
    if not checkboxes:
        raise GestorError(
            "Aviso de seguranca RemoteApp",
            "As opcoes de recursos nao ficaram disponiveis no tempo esperado.",
        )
    marcadas = set()
    sem_novidades = 0
    for rodada in range(6):
        quantidade_antes = len(marcadas)
        try:
            checkboxes = janela_seg.descendants(control_type="CheckBox")
        except Exception:
            checkboxes = []
        for caixa in checkboxes:
            try:
                if not caixa.is_visible():
                    continue
                nome = caixa.window_text() or f"checkbox-{caixa.handle}"
                estado = caixa.get_toggle_state()
                if estado != 1:
                    try:
                        caixa.toggle()
                    except Exception:
                        caixa.click_input()
                    time.sleep(0.35)
                if caixa.get_toggle_state() != 1:
                    try:
                        caixa.scroll_into_view()
                        try:
                            caixa.toggle()
                        except Exception:
                            caixa.click_input()
                        time.sleep(0.35)
                    except Exception:
                        pass
                if caixa.get_toggle_state() != 1:
                    raise GestorError(
                        "Aviso de seguranca RemoteApp",
                        f"Nao foi possivel marcar: {nome}",
                    )
                marcadas.add(nome)
            except GestorError:
                raise
            except Exception:
                continue
        sem_novidades = sem_novidades + 1 if len(marcadas) == quantidade_antes else 0
        if rodada < 5 and sem_novidades < 2:
            import pyautogui

            left, top, right, bottom = _retangulo_visivel(seguranca["hwnd"])
            pyautogui.moveTo(left + int((right - left) * 0.55), top + int((bottom - top) * 0.69))
            pyautogui.scroll(-6)
            time.sleep(0.6)
        else:
            break
    if not marcadas:
        raise GestorError(
            "Aviso de seguranca RemoteApp", "Nenhuma checkbox foi localizada."
        )
    janelas_antes_menu = _estado_janelas()
    inicio_conexao = time.monotonic()
    _invocar_controle(_controle_texto(janela_seg, "Conectar", tipos=("Button",)))
    log_gestor(log, f"RemoteApp conectado; {len(marcadas)} recursos permitidos.")
    menu = aguardar_janela(
        ("remote", "app"),
        config.timeout_janela,
        novos_desde=janelas_antes_menu,
        classe="RAIL_WINDOW",
        etapa="Menu Remote App",
    )
    restante = 10 - (time.monotonic() - inicio_conexao)
    if restante > 0:
        time.sleep(restante)
    return menu


def abrir_lojas_hoje(config, log, menu=None):
    if menu is None:
        menu = aguardar_janela(
            ("remote", "app"),
            config.timeout_janela,
            classe="RAIL_WINDOW",
            etapa="Menu Remote App",
        )
    preparar_janela_gestor(menu["hwnd"], log)
    anteriores = _estado_janelas()
    _clicar_texto_ocr_janela(menu["hwnd"], "Lojas Hoje", log)
    gestor = aguardar_janela(
        ("gestor", "erp"),
        config.timeout_janela,
        novos_desde=anteriores,
        etapa="Gestor ERP",
    )
    preparar_janela_gestor(gestor["hwnd"], log)
    log_gestor(log, "Lojas Hoje aberta no Gestor ERP.")
    time.sleep(3)
    return gestor


def _formulario_acesso_visivel(palavras):
    texto = normalizar(" ".join(str(item.get("text", "")) for item in palavras))
    return all(termo in texto for termo in ("usuario", "senha", "acessar"))


def _ler_ocr_monitor_principal(caminho):
    area = _monitor_principal()["area_trabalho"]
    ImageGrab.grab(bbox=area, all_screens=True).save(caminho)
    return _executar_ocr(caminho)


def _aguardar_formulario_acesso_pronto(hwnd, timeout, log):
    fim = time.monotonic() + timeout
    inicio = time.monotonic()
    log_gestor(log, "Aguardando 8 segundos para o formulario remoto terminar de carregar...")
    while time.monotonic() < fim:
        if not _janela_existe(hwnd):
            raise GestorError(
                "Controle de acesso",
                "A janela de acesso desapareceu durante a espera de carregamento.",
            )
        if time.monotonic() - inicio >= 8:
            log_gestor(log, "Espera passiva concluida; formulario remoto permaneceu aberto.")
            return
        time.sleep(0.5)
    raise GestorError(
        "Controle de acesso",
        "O formulario nao permaneceu disponivel pelo tempo minimo de carregamento.",
    )


def _aguardar_usuario_acesso_digitado(hwnd, usuario, timeout=12):
    temporario = Path(tempfile.gettempdir()) / f"gestor_usuario_{os.getpid()}.png"
    fim = time.monotonic() + timeout
    try:
        while time.monotonic() < fim:
            area = _monitor_principal()["area_trabalho"]
            imagem = ImageGrab.grab(
                bbox=(area[0] + 105, area[1] + 215, area[0] + 500, area[1] + 252),
                all_screens=True,
            )
            imagem = ImageOps.autocontrast(ImageOps.grayscale(imagem))
            imagem = imagem.resize(
                (imagem.width * 6, imagem.height * 6), Image.Resampling.LANCZOS
            )
            imagem.save(temporario)
            palavras = _executar_ocr(temporario)
            texto = normalizar(" ".join(str(item.get("text", "")) for item in palavras))
            if normalizar(usuario) in texto:
                return
            time.sleep(0.8)
    finally:
        temporario.unlink(missing_ok=True)
    raise GestorError(
        "Controle de acesso",
        "O usuario digitado nao foi confirmado visualmente; o acesso nao foi enviado.",
    )


def _ler_clipboard_windows():
    user32 = ctypes.windll.user32
    kernel32 = ctypes.windll.kernel32
    user32.GetClipboardData.argtypes = [wintypes.UINT]
    user32.GetClipboardData.restype = wintypes.HANDLE
    kernel32.GlobalLock.argtypes = [wintypes.HANDLE]
    kernel32.GlobalLock.restype = ctypes.c_void_p
    kernel32.GlobalUnlock.argtypes = [wintypes.HANDLE]
    if not user32.OpenClipboard(None):
        raise GestorError("Controle de acesso", "Nao foi possivel abrir a area de transferencia.")
    try:
        handle = user32.GetClipboardData(13)  # CF_UNICODETEXT
        if not handle:
            return ""
        ponteiro = kernel32.GlobalLock(handle)
        if not ponteiro:
            return ""
        try:
            return ctypes.wstring_at(ponteiro)
        finally:
            kernel32.GlobalUnlock(handle)
    finally:
        user32.CloseClipboard()


def _limpar_clipboard_windows():
    user32 = ctypes.windll.user32
    for _tentativa in range(5):
        if user32.OpenClipboard(None):
            try:
                user32.EmptyClipboard()
                return
            finally:
                user32.CloseClipboard()
        time.sleep(0.1)
    raise GestorError("Area de transferencia", "Nao foi possivel limpar o clipboard.")


def _confirmar_campo_ativo(valor_esperado, etapa="Controle de acesso", nome="usuario"):
    import pyautogui

    _limpar_clipboard_windows()
    pyautogui.hotkey("ctrl", "a")
    time.sleep(0.2)
    pyautogui.hotkey("ctrl", "c")
    copiado = ""
    fim = time.monotonic() + 3
    while time.monotonic() < fim:
        time.sleep(0.2)
        copiado = _ler_clipboard_windows()
        if copiado:
            break
    if copiado.strip() != str(valor_esperado).strip():
        raise GestorError(
            etapa,
            f"O campo {nome} nao recebeu o valor esperado; a etapa nao foi confirmada.",
        )


def _ponto_canvas_acesso(hwnd, x, y):
    left, top, right, bottom = _retangulo_visivel(hwnd)
    return left + int((right - left) * x), top + int((bottom - top) * y)


def _clicar_retangulo_real(hwnd, x, y, cliques=1):
    import pyautogui

    _ativar(hwnd)
    pyautogui.click(*_ponto_canvas_acesso(hwnd, x, y), clicks=cliques, interval=0.15)


def _preencher_campo_canvas_acesso(hwnd, x, y, texto, intervalo):
    import pyautogui

    _ativar(hwnd)
    px, py = _ponto_canvas_acesso(hwnd, x, y)
    pyautogui.click(px, py, clicks=2, interval=0.15)
    time.sleep(0.8)
    pyautogui.hotkey("ctrl", "a")
    time.sleep(0.25)
    pyautogui.press("backspace")
    time.sleep(0.25)
    pyautogui.write(str(texto), interval=intervalo)


def _data_campo_corresponde(texto, data_esperada):
    esperado = data_esperada.strftime("%d%m%Y")
    dia, mes, ano = esperado[:2], esperado[2:4], esperado[4:]
    aceitos = {
        esperado,
        f"{dia}1{mes}{ano}",
        f"{dia}{mes}1{ano}",
        f"{dia}1{mes}1{ano}",
    }
    candidatos = re.split(r"\s*\|\s*", str(texto or ""))
    return any(re.sub(r"\D", "", item) in aceitos for item in candidatos)


def _ler_data_visual_campo(hwnd, x, y):
    _ativar(hwnd)
    time.sleep(0.3)
    left, top, right, bottom = _retangulo_visivel(hwnd)
    largura = right - left
    altura = bottom - top
    centro_x = left + int(largura * x)
    centro_y = top + int(altura * y)
    meia_largura = max(44, int(largura * 0.036))
    meia_altura = max(13, int(altura * 0.018))
    imagem = ImageGrab.grab(
        bbox=(
            centro_x - meia_largura,
            centro_y - meia_altura,
            centro_x + meia_largura,
            centro_y + meia_altura,
        ),
        all_screens=True,
    )
    cinza = ImageOps.autocontrast(ImageOps.grayscale(imagem))
    variantes = (
        imagem.resize((imagem.width * 8, imagem.height * 8), Image.Resampling.LANCZOS),
        cinza.resize((imagem.width * 8, imagem.height * 8), Image.Resampling.LANCZOS),
        cinza.point(lambda pixel: 255 if pixel > 180 else 0).resize(
            (imagem.width * 8, imagem.height * 8), Image.Resampling.NEAREST
        ),
    )
    textos = []
    for indice, variante in enumerate(variantes, start=1):
        temporario = Path(tempfile.gettempdir()) / (
            f"gestor_data_{os.getpid()}_{int(x * 1000)}_{int(y * 1000)}_{indice}.png"
        )
        try:
            variante.save(temporario)
            palavras = _executar_ocr(temporario)
            texto = " ".join(
                str(item.get("text", "")) for item in palavras
            ).strip()
            if texto:
                textos.append(texto)
        finally:
            temporario.unlink(missing_ok=True)
    return " | ".join(textos)


def _preencher_validar_data_canvas(hwnd, x, y, data, intervalo, nome):
    import pyautogui

    texto = data.strftime("%d/%m/%Y")
    ultimo = ""
    for tentativa in range(1, 3):
        _preencher_campo_canvas_acesso(hwnd, x, y, texto, intervalo)
        pyautogui.press("tab")
        time.sleep(0.8)
        ultimo = _ler_data_visual_campo(hwnd, x, y)
        if _data_campo_corresponde(ultimo, data):
            return ultimo
        if tentativa == 1:
            time.sleep(0.5)
    raise GestorError(
        "Periodo do relatorio",
        f"{nome} nao foi confirmada. Esperado={texto}; lido={ultimo!r}.",
        hwnd=hwnd,
    )


def _clicar_canvas_acesso(hwnd, x, y):
    _clicar_retangulo_real(hwnd, x, y)


def _capturar_campo_canvas_acesso(hwnd, y_inicio, y_fim):
    left, top, right, bottom = _retangulo_visivel(hwnd)
    largura = right - left
    altura = bottom - top
    return ImageGrab.grab(
        bbox=(
            left + int(largura * 0.13),
            top + int(altura * y_inicio),
            left + int(largura * 0.54),
            top + int(altura * y_fim),
        ),
        all_screens=True,
    )


def _validar_mudanca_visual_campo(antes, depois, nome):
    if ImageChops.difference(antes.convert("RGB"), depois.convert("RGB")).getbbox() is None:
        raise GestorError(
            "Controle de acesso",
            f"O campo {nome} nao apresentou mudanca visual depois da digitacao.",
        )


def autenticar_comercial(gestor, config, log):
    anteriores = _estado_janelas()
    _clicar_modulo_comercial(gestor["hwnd"], log)
    try:
        acesso = aguardar_janela(
            ("controle", "acesso"),
            min(12, config.timeout_janela),
            novos_desde=anteriores,
            etapa="Controle de acesso",
        )
    except GestorError:
        log_gestor(
            log,
            "O primeiro clique em Comercial nao gerou janela; repetindo uma vez no mesmo icone.",
        )
        _clicar_modulo_comercial(gestor["hwnd"], log)
        acesso = aguardar_janela(
            ("controle", "acesso"),
            config.timeout_janela,
            novos_desde=anteriores,
            etapa="Controle de acesso",
        )
    preparar_janela_gestor(acesso["hwnd"], log)

    hwnd_acesso = acesso["hwnd"]
    _aguardar_formulario_acesso_pronto(hwnd_acesso, config.timeout_janela, log)
    usuario_antes = _capturar_campo_canvas_acesso(hwnd_acesso, 0.33, 0.38)
    _preencher_campo_canvas_acesso(
        hwnd_acesso, 0.33, 0.35, config.usuario.upper(), config.intervalo_tecla
    )
    time.sleep(1)
    usuario_depois = _capturar_campo_canvas_acesso(hwnd_acesso, 0.33, 0.38)
    _validar_mudanca_visual_campo(usuario_antes, usuario_depois, "Usuario")
    log_gestor(log, "Campo Usuario apresentou mudanca visual apos a digitacao.")

    senha_antes = _capturar_campo_canvas_acesso(hwnd_acesso, 0.38, 0.43)
    _preencher_campo_canvas_acesso(
        hwnd_acesso, 0.33, 0.40, config.senha.upper(), config.intervalo_tecla
    )
    time.sleep(1)
    senha_depois = _capturar_campo_canvas_acesso(hwnd_acesso, 0.38, 0.43)
    _validar_mudanca_visual_campo(senha_antes, senha_depois, "Senha")
    log_gestor(log, "Campo Senha apresentou mudanca visual e permanece mascarado.")

    _clicar_canvas_acesso(hwnd_acesso, 0.86, 0.94)
    log_gestor(log, "Credenciais enviadas; aguardando carregar a lista de empresas.")
    time.sleep(3)
    if _janela_existe(hwnd_acesso):
        _clicar_canvas_acesso(hwnd_acesso, 0.45, 0.50)
        time.sleep(0.7)
        _clicar_canvas_acesso(hwnd_acesso, 0.86, 0.94)
        log_gestor(log, "Empresa Lojas Hoje selecionada e segundo Acessar enviado.")
    _aguardar_fechar(hwnd_acesso, config.timeout_janela)
    time.sleep(2)
    log_gestor(log, "Controle de acesso autenticado.")

    anteriores = _estado_janelas()
    _clicar_modulo_comercial(gestor["hwnd"], log)
    try:
        local = aguardar_janela(
            ("selecionar", "local", "trabalho"),
            min(12, config.timeout_janela),
            novos_desde=anteriores,
            etapa="Selecionar local de trabalho",
        )
    except GestorError:
        log_gestor(
            log,
            "O primeiro clique apos o login nao abriu os locais; repetindo uma vez no mesmo icone Comercial.",
        )
        _clicar_modulo_comercial(gestor["hwnd"], log)
        local = aguardar_janela(
            ("selecionar", "local", "trabalho"),
            config.timeout_janela,
            novos_desde=anteriores,
            etapa="Selecionar local de trabalho",
        )
    return local


def selecionar_local_99(local, config, log):
    preparar_janela_gestor(local["hwnd"], log)
    _clicar_retangulo_real(local["hwnd"], 0.10, 0.315)
    time.sleep(0.6)
    anteriores = _estado_janelas()
    _clicar_retangulo_real(local["hwnd"], 0.685, 0.93)
    log_gestor(log, "Local 99 marcado e botao Selecionar o Local acionado.")
    impressora = aguardar_janela(
        ("selecionar", "impressora", "padrao"),
        config.timeout_janela,
        novos_desde=anteriores,
        etapa="Impressora padrao",
    )
    log_gestor(log, "Local 99 selecionado.")
    return impressora


def selecionar_impressora_windows(impressora, config, log):
    preparar_janela_gestor(impressora["hwnd"], log)
    _ativar(impressora["hwnd"])
    anteriores = _estado_janelas()
    _clicar_retangulo_real(impressora["hwnd"], 0.098, 0.736)
    log_gestor(log, "Opcao Impressora Padrao Windows acionada.")
    modulo = aguardar_janela(
        "gestor",
        config.timeout_janela,
        novos_desde=anteriores,
        etapa="Modulo Comercial",
    )
    preparar_janela_gestor(modulo["hwnd"], log)
    log_gestor(log, "Impressora Padrao Windows selecionada.")
    return modulo


def abrir_metas_vendas_local(modulo, config, log):
    _clicar(modulo["hwnd"], 0.068, 0.035)
    time.sleep(0.6)
    _mover(modulo["hwnd"], 0.115, 0.085)
    time.sleep(0.8)
    anteriores = _estado_janelas()
    _clicar(modulo["hwnd"], 0.271, 0.086)
    liberacao = aguardar_janela(
        ("liberacao", "acesso"),
        config.timeout_janela,
        novos_desde=anteriores,
        etapa="Liberacao de acesso",
    )
    return liberacao


def autenticar_liberacao(liberacao, config, log):
    preparar_janela_gestor(liberacao["hwnd"], log)
    _ativar(liberacao["hwnd"])
    time.sleep(1.5)
    _preencher_campo_canvas_acesso(
        liberacao["hwnd"],
        0.52,
        0.31,
        config.usuario.upper(),
        config.intervalo_tecla,
    )
    time.sleep(0.5)
    _preencher_campo_canvas_acesso(
        liberacao["hwnd"],
        0.52,
        0.50,
        config.senha.upper(),
        config.intervalo_tecla,
    )
    time.sleep(0.8)
    anteriores = _estado_janelas()
    _clicar_retangulo_real(liberacao["hwnd"], 0.54, 0.89)
    log_gestor(log, "Credenciais da Liberacao de Acesso preenchidas; Confirmar acionado.")
    metas = aguardar_janela(
        ("metas", "vendas", "local"),
        config.timeout_janela,
        novos_desde=anteriores,
        etapa="Metas de vendas por local",
    )
    preparar_janela_gestor(metas["hwnd"], log)
    log_gestor(log, "Liberacao de acesso confirmada.")
    return metas


def gerar_relatorio_metas(metas, config, log, agora=None):
    agora = agora or datetime.now()
    data = agora.strftime("%d/%m/%Y")
    preparar_janela_gestor(metas["hwnd"], log)
    _ativar(metas["hwnd"])
    _clicar_texto_ocr_janela(
        metas["hwnd"], "Relatorio", log, faixa_y=(0.75, 0.99)
    )
    log_gestor(log, "Secao Opcoes p/Relatorio ativada pelo botao Relatorio inferior.")
    time.sleep(0.7)
    alvo_iniciar = _localizar_texto_ocr_janela(
        metas["hwnd"], "Iniciar", faixa_y=(0.75, 0.99)
    )
    left, top, right, bottom = _retangulo_visivel(metas["hwnd"])
    deslocamento_periodo = max(55, alvo_iniciar["altura"] * 5)
    y_periodo = (alvo_iniciar["y"] - deslocamento_periodo - top) / (bottom - top)
    if not 0.75 <= y_periodo <= 0.92:
        raise GestorError(
            "Periodo do relatorio",
            f"Linha inferior das datas calculada fora da faixa segura: y={y_periodo:.3f}.",
            hwnd=metas["hwnd"],
        )
    log_gestor(
        log,
        f"Campos de Periodo localizados pela ancora Iniciar: y={y_periodo:.3f}.",
    )
    inicial_lida = _preencher_validar_data_canvas(
        metas["hwnd"], 0.228, y_periodo, agora, config.intervalo_tecla, "Data inicial"
    )
    time.sleep(0.5)
    final_lida = _preencher_validar_data_canvas(
        metas["hwnd"], 0.312, y_periodo, agora, config.intervalo_tecla, "Data final"
    )
    # Releia o primeiro campo depois do segundo. Isso detecta o atraso de foco
    # que anteriormente enviava as duas datas para o mesmo controle.
    inicial_revalidada = _ler_data_visual_campo(metas["hwnd"], 0.228, y_periodo)
    if not _data_campo_corresponde(inicial_revalidada, agora):
        raise GestorError(
            "Periodo do relatorio",
            "A Data inicial mudou depois do preenchimento da Data final: "
            f"{inicial_revalidada!r}.",
            hwnd=metas["hwnd"],
        )
    log_gestor(
        log,
        "Periodo inferior validado: "
        f"inicial={inicial_lida!r}; final={final_lida!r}; data={data}.",
    )
    time.sleep(0.8)
    anteriores = _estado_janelas()
    _clicar_texto_ocr_janela(
        metas["hwnd"], "Iniciar", log, faixa_y=(0.75, 0.99)
    )
    log_gestor(log, "Botao Iniciar inferior acionado.")
    try:
        impressao = aguardar_janela(
            ("impressao", "relatorio"),
            min(12, config.timeout_janela),
            novos_desde=anteriores,
            etapa="Impressao de relatorio",
        )
    except GestorError:
        if not _janela_existe(metas["hwnd"]):
            raise GestorError(
                "Impressao de relatorio",
                "A tela de metas desapareceu e a janela de impressao nao abriu.",
                hwnd=metas["hwnd"],
            )
        log_gestor(
            log,
            "A janela de impressao nao abriu; datas ainda validas e tela anterior "
            "presente. Repetindo Iniciar uma unica vez.",
        )
        _clicar_texto_ocr_janela(
            metas["hwnd"], "Iniciar", log, faixa_y=(0.75, 0.99)
        )
        impressao = aguardar_janela(
            ("impressao", "relatorio"),
            config.timeout_janela,
            novos_desde=anteriores,
            etapa="Impressao de relatorio",
        )
    preparar_janela_gestor(impressao["hwnd"], log)
    anteriores = _estado_janelas()
    _clicar_retangulo_real(impressao["hwnd"], 0.07, 0.31)
    log_gestor(log, "Opcao Visualizar Relatorio em Tela selecionada.")
    time.sleep(0.7)
    _clicar_retangulo_real(impressao["hwnd"], 0.405, 0.915)
    log_gestor(log, "Botao OK da janela Impressao de Relatorio acionado.")
    time.sleep(4)
    try:
        preview = aguardar_janela(
            ("report", "preview"),
            min(12, config.timeout_relatorio),
            novos_desde=anteriores,
            etapa="Report Preview",
        )
    except GestorError:
        if not _janela_existe(impressao["hwnd"]):
            raise GestorError(
                "Report Preview",
                "A janela de impressao fechou, mas o Report Preview nao apareceu.",
                hwnd=metas["hwnd"],
            )
        log_gestor(
            log,
            "Report Preview ainda nao abriu e a janela de impressao permanece "
            "visivel. Repetindo OK uma unica vez.",
        )
        _clicar_retangulo_real(impressao["hwnd"], 0.405, 0.915)
        time.sleep(4)
        try:
            preview = aguardar_janela(
                ("report", "preview"),
                config.timeout_relatorio,
                novos_desde=anteriores,
                etapa="Report Preview",
            )
        except GestorError as erro:
            raise GestorError(
                "Report Preview",
                str(erro),
                hwnd=impressao["hwnd"],
            ) from erro
    preparar_janela_gestor(preview["hwnd"], log)
    log_gestor(log, f"Relatorio solicitado para {data} ate {data}.")
    return preview, data


def _capturar_janela(hwnd, caminho):
    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    bbox = _retangulo_visivel(hwnd)
    imagem = ImageGrab.grab(bbox=bbox, all_screens=True)
    imagem.save(caminho)
    return caminho


def _selecionar_zoom_preview(hwnd, percentual):
    try:
        janela = _janela_pywinauto(hwnd, backend="win32")
        combos = janela.descendants(class_name="TComboBox")
        if not combos:
            combos = janela.descendants(class_name="ComboBox")
        for combo in combos:
            try:
                itens = combo.item_texts()
            except Exception:
                itens = []
            alvo = next(
                (item for item in itens if normalizar(item) == normalizar(percentual)),
                None,
            )
            if alvo:
                combo.select(alvo)
                time.sleep(1.2)
                return True
    except Exception:
        return False
    return False


def capturar_report_preview(preview, config, log):
    config.pasta_logs.mkdir(parents=True, exist_ok=True)
    pasta = config.pasta_logs / "evidencias"
    pasta.mkdir(parents=True, exist_ok=True)
    nome = f"metas_realizado_{datetime.now():%d-%m-%Y_%H-%M-%S}.png"
    caminho = pasta / nome
    _selecionar_zoom_preview(preview["hwnd"], "100%")
    _capturar_janela(preview["hwnd"], caminho)
    log_gestor(log, f"Evidencia do relatorio salva: {caminho}")
    return caminho


def _executar_ocr(caminho):
    script = Path(__file__).resolve().with_name("ocr_windows.ps1")
    processo = subprocess.run(
        [
            "powershell.exe",
            "-NoProfile",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(script),
            str(Path(caminho).resolve()),
        ],
        capture_output=True,
        check=False,
        timeout=90,
    )
    stdout = processo.stdout.decode("utf-8-sig", errors="replace").strip()
    stderr = processo.stderr.decode("utf-8-sig", errors="replace").strip()
    if processo.returncode != 0:
        raise GestorError("OCR", stderr or "Windows OCR terminou com erro.")
    if not stdout:
        return []
    dados = json.loads(stdout)
    if isinstance(dados, dict):
        dados = [dados]
    return dados


def _agrupar_linhas_ocr(palavras, tolerancia=14):
    ordenadas = sorted(
        palavras,
        key=lambda p: (int(p.get("y", 0)) + int(p.get("height", 0)) // 2, int(p.get("x", 0))),
    )
    linhas = []
    for palavra in ordenadas:
        centro = int(palavra.get("y", 0)) + int(palavra.get("height", 0)) // 2
        destino = None
        for linha in linhas:
            if abs(linha["centro"] - centro) <= tolerancia:
                destino = linha
                break
        if destino is None:
            destino = {"centro": centro, "palavras": []}
            linhas.append(destino)
        destino["palavras"].append(palavra)
        centros = [
            int(p.get("y", 0)) + int(p.get("height", 0)) // 2
            for p in destino["palavras"]
        ]
        destino["centro"] = sum(centros) // len(centros)
    for linha in linhas:
        linha["palavras"].sort(key=lambda p: int(p.get("x", 0)))
        linha["texto"] = " ".join(str(p.get("text", "")) for p in linha["palavras"])
    return sorted(linhas, key=lambda linha: linha["centro"])


def _palavras_do_alvo_ocr(linha, texto):
    """Retorna somente as palavras que formam o alvo, nunca a linha inteira."""
    palavras = linha.get("palavras", [])
    alvo = normalizar(texto)
    alvo_compacto = alvo.replace(" ", "")
    candidatos = []

    def distancia_edicao(a, b):
        anterior = list(range(len(b) + 1))
        for indice_a, caractere_a in enumerate(a, start=1):
            atual = [indice_a]
            for indice_b, caractere_b in enumerate(b, start=1):
                atual.append(
                    min(
                        atual[-1] + 1,
                        anterior[indice_b] + 1,
                        anterior[indice_b - 1] + (caractere_a != caractere_b),
                    )
                )
            anterior = atual
        return anterior[-1]

    for inicio in range(len(palavras)):
        for fim in range(inicio + 1, len(palavras) + 1):
            trecho = normalizar(" ".join(str(p.get("text", "")) for p in palavras[inicio:fim]))
            trecho_compacto = trecho.replace(" ", "")
            exato = trecho == alvo or trecho_compacto == alvo_compacto
            toleravel = (
                len(alvo_compacto) >= 6
                and abs(len(trecho_compacto) - len(alvo_compacto)) <= 1
                and distancia_edicao(trecho_compacto, alvo_compacto) <= 1
            )
            if exato or toleravel:
                candidatos.append(palavras[inicio:fim])
    if not candidatos:
        return None
    return min(candidatos, key=len)


def _localizar_texto_ocr_janela(
    hwnd,
    texto,
    timeout=30,
    faixa_x=None,
    faixa_y=None,
):
    temporario = Path(tempfile.gettempdir()) / f"gestor_menu_{os.getpid()}.png"
    fim = time.monotonic() + timeout
    ultimas_linhas = []
    try:
        while time.monotonic() < fim:
            _ativar(hwnd)
            retangulo = _retangulo_visivel(hwnd)
            _capturar_janela(hwnd, temporario)
            linhas = _agrupar_linhas_ocr(_executar_ocr(temporario), tolerancia=10)
            ultimas_linhas = [item["texto"] for item in linhas]
            resultados = []
            for linha in linhas:
                palavras = _palavras_do_alvo_ocr(linha, texto)
                if palavras is None:
                    continue
                esquerda = min(int(item["x"]) for item in palavras)
                direita = max(int(item["x"]) + int(item["width"]) for item in palavras)
                topo = min(int(item["y"]) for item in palavras)
                base = max(int(item["y"]) + int(item["height"]) for item in palavras)
                ancora_x = retangulo[0] + (esquerda + direita) // 2
                ancora_y = retangulo[1] + (topo + base) // 2
                if faixa_x is not None:
                    monitor = _monitor_principal()["retangulo"]
                    largura_monitor = monitor[2] - monitor[0]
                    minimo_x = monitor[0] + int(largura_monitor * faixa_x[0])
                    maximo_x = monitor[0] + int(largura_monitor * faixa_x[1])
                    if not minimo_x <= ancora_x <= maximo_x:
                        continue
                if faixa_y is not None:
                    altura = retangulo[3] - retangulo[1]
                    minimo_y = retangulo[1] + int(altura * faixa_y[0])
                    maximo_y = retangulo[1] + int(altura * faixa_y[1])
                    if not minimo_y <= ancora_y <= maximo_y:
                        continue
                texto_reconhecido = " ".join(str(item.get("text", "")) for item in palavras)
                resultados.append(
                    {
                        "x": ancora_x,
                        "y": ancora_y,
                        "largura": direita - esquerda,
                        "altura": base - topo,
                        "texto_ocr": texto_reconhecido,
                        "linha": linha["texto"],
                    }
                )
            if resultados:
                return resultados[0]
            time.sleep(1)
        raise GestorError(
            "Interface visual",
            f"Opcao visual nao encontrada: {texto}. OCR={ultimas_linhas}",
            hwnd=hwnd,
        )
    finally:
        temporario.unlink(missing_ok=True)


def _clicar_texto_ocr_janela(
    hwnd,
    texto,
    log,
    cliques=1,
    timeout=30,
    deslocamento_x=0,
    deslocamento_y=0,
    faixa_x=None,
    faixa_y=None,
):
    import pyautogui

    alvo = _localizar_texto_ocr_janela(
        hwnd,
        texto,
        timeout=timeout,
        faixa_x=faixa_x,
        faixa_y=faixa_y,
    )
    x = alvo["x"] + deslocamento_x
    y = alvo["y"] + deslocamento_y
    log_gestor(
        log,
        f"Clicando visualmente no alvo exato {texto!r} "
        f"(OCR={alvo['texto_ocr']!r}; ancora=({alvo['x']}, {alvo['y']}); "
        f"deslocamento=({deslocamento_x}, {deslocamento_y})): ({x}, {y}).",
    )
    _ativar(hwnd)
    pyautogui.click(x, y, clicks=cliques, interval=0.15)
    return {**alvo, "x": x, "y": y}


def _clicar_modulo_comercial(hwnd, log):
    # O texto serve apenas como ancora. O controle clicavel e o quadrado azul
    # imediatamente acima dele, que e o primeiro modulo da esquerda.
    log_gestor(log, "Selecionando o primeiro modulo da esquerda: Comercial.")
    alvo = _clicar_texto_ocr_janela(
        hwnd,
        "Comercial",
        log,
        cliques=2,
        deslocamento_y=-90,
        faixa_x=(0.15, 0.40),
    )
    log_gestor(log, "Duplo clique confirmado no icone do primeiro modulo: Comercial.")
    return alvo


def _corrigir_token_monetario(token):
    token = token.replace(" ", "").replace("'", "").replace("`", "")
    token = token.translate(
        str.maketrans(
            {
                "O": "0",
                "o": "0",
                "Q": "0",
                "D": "0",
                "C": "0",
                "I": "1",
                "l": "1",
                "|": "1",
                "S": "5",
                "s": "5",
                "B": "8",
            }
        )
    )
    return token


def _valores_monetarios_linha(texto):
    candidatos = re.findall(
        r"\d[\d\.\sOQoDCIl|SsB]*[,;]\s*[\dOQoDCIl|SsB]{2}(?![\dOQoDCIl|SsB])",
        texto,
    )
    valores = []
    for candidato in candidatos:
        corrigido = _corrigir_token_monetario(candidato).replace(";", ",")
        try:
            valores.append(numero_brasileiro(corrigido))
        except ValueError:
            continue
    return valores


def extrair_valores_realizados(caminho_imagem):
    palavras = _executar_ocr(caminho_imagem)
    linhas = _agrupar_linhas_ocr(palavras)
    valores = {}
    total = None
    textos = []
    for linha in linhas:
        texto = linha["texto"]
        textos.append(texto)
        texto_normalizado = normalizar(texto)
        monetarios = _valores_monetarios_linha(texto)
        for loja in LOJAS:
            if re.search(rf"(?:^|[^a-z]){loja.casefold()}(?:$|[^a-z])", texto_normalizado):
                if len(monetarios) >= 2:
                    if loja in valores:
                        raise GestorError("OCR", f"Loja {loja} apareceu mais de uma vez.")
                    valores[loja] = monetarios[1]
        if "total geral" in texto_normalizado and len(monetarios) >= 2:
            total = monetarios[1]
    if set(valores) != set(LOJAS) or total is None:
        encontrados = {loja: formatar_brasileiro(v) for loja, v in valores.items()}
        raise GestorError(
            "OCR",
            "Leitura incompleta do Valor Realizado. "
            f"Lojas={encontrados}; total={total}; linhas={textos}",
        )
    validar_total_relatorio(valores, total)
    return valores, total, linhas


def _imagem_ampliada(caminho_origem, caminho_destino, escala=2):
    with Image.open(caminho_origem) as imagem:
        cinza = ImageOps.grayscale(imagem)
        ampliada = cinza.resize(
            (cinza.width * escala, cinza.height * escala), Image.Resampling.LANCZOS
        )
        ampliada.save(caminho_destino)
    return caminho_destino


def extrair_relatorio_com_tentativas(preview, evidencia, config, log):
    erros = []
    tentativas = [("100%", evidencia), ("75%", None), ("125%", None)]
    for indice, (zoom, caminho) in enumerate(tentativas, start=1):
        if caminho is None:
            _selecionar_zoom_preview(preview["hwnd"], zoom)
            caminho = config.pasta_erros / f"gestor_ocr_{zoom.replace('%', '')}_{datetime.now():%Y%m%d_%H%M%S}.png"
            _capturar_janela(preview["hwnd"], caminho)
        ampliada = Path(tempfile.gettempdir()) / f"gestor_ocr_ampliada_{os.getpid()}_{indice}.png"
        try:
            _imagem_ampliada(caminho, ampliada, escala=2)
            valores, total, linhas = extrair_valores_realizados(ampliada)
            log_gestor(log, f"OCR validado com zoom {zoom}.")
            return valores, total, linhas
        except Exception as erro:
            erros.append(f"{zoom}: {erro}")
        finally:
            try:
                ampliada.unlink(missing_ok=True)
            except OSError:
                pass
    raise GestorError("OCR", " | ".join(erros))


def _arquivo_liberado(arquivo):
    try:
        with Path(arquivo).open("r+b") as handle:
            if os.name == "nt":
                import msvcrt

                msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
                handle.seek(0)
                msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
        return True
    except OSError:
        return False


def aguardar_excel_liberado(arquivo, timeout=120):
    fim = time.monotonic() + timeout
    while time.monotonic() < fim:
        if _arquivo_liberado(arquivo):
            return
        time.sleep(1)
    raise GestorError("Excel", f"Arquivo continua bloqueado: {arquivo}")


def _resolver_referencia_simples(ws, coordenada):
    visitadas = set()
    atual = coordenada
    for _ in range(20):
        if atual in visitadas:
            raise GestorError("Excel", f"Referencia circular encontrada em {coordenada}.")
        visitadas.add(atual)
        valor = ws[atual].value
        if not isinstance(valor, str) or not valor.startswith("="):
            return valor
        referencia = re.fullmatch(r"=\$?([A-Z]{1,3})\$?(\d+)", valor.strip(), re.I)
        if not referencia:
            return valor
        atual = f"{referencia.group(1).upper()}{referencia.group(2)}"
    raise GestorError("Excel", f"Cadeia de referencias muito longa em {coordenada}.")


def _validar_estrutura_excel(ws):
    cabecalho = normalizar(ws["C18"].value)
    if "venda dia" not in cabecalho:
        raise GestorError(
            "Excel", f"C18 nao corresponde a Venda dia: {ws['C18'].value!r}"
        )
    for loja, linha in LINHAS_EXCEL.items():
        rotulos = " ".join(
            str(_resolver_referencia_simples(ws, ws.cell(linha, coluna).coordinate) or "")
            for coluna in (1, 2)
        )
        if not re.search(rf"(?:^|[^a-z]){loja.casefold()}(?:$|[^a-z])", normalizar(rotulos)):
            raise GestorError(
                "Excel", f"Linha {linha} nao corresponde a {loja}: {rotulos!r}"
            )


def validar_excel_metas(arquivo, valores_esperados=None):
    arquivo = Path(arquivo)
    if not arquivo.exists() or arquivo.stat().st_size <= 0:
        raise GestorError("Excel", f"Arquivo invalido: {arquivo}")
    wb = load_workbook(arquivo, data_only=False, keep_links=True)
    try:
        if "base" not in wb.sheetnames:
            raise GestorError("Excel", "Aba 'base' nao encontrada.")
        ws = wb["base"]
        _validar_estrutura_excel(ws)
        if valores_esperados is not None:
            for loja, linha in LINHAS_EXCEL.items():
                atual = ws[f"C{linha}"].value
                if atual is None:
                    raise GestorError("Excel", f"C{linha} ficou vazia.")
                if Decimal(str(atual)).quantize(Decimal("0.01")) != Decimal(
                    str(valores_esperados[loja])
                ).quantize(Decimal("0.01")):
                    raise GestorError(
                        "Excel",
                        f"C{linha} diverge de {loja}: {atual!r}",
                    )
        return {loja: ws[f"C{linha}"].value for loja, linha in LINHAS_EXCEL.items()}
    finally:
        wb.close()


def atualizar_excel_metas(arquivo, valores, pasta_backup, log=None):
    arquivo = Path(arquivo)
    pasta_backup = Path(pasta_backup)
    soma = sum(
        (Decimal(str(valores[loja])) for loja in LOJAS),
        Decimal("0"),
    )
    validar_total_relatorio(valores, soma)
    aguardar_excel_liberado(arquivo)
    validar_excel_metas(arquivo)
    pasta_backup.mkdir(parents=True, exist_ok=True)
    backup = pasta_backup / f"{arquivo.stem}_antes_gestor_{datetime.now():%Y-%m-%d_%H-%M-%S}{arquivo.suffix}"
    shutil.copy2(arquivo, backup)
    temporario = arquivo.with_name(f".{arquivo.stem}.gestor-{os.getpid()}.tmp{arquivo.suffix}")
    wb = None
    try:
        wb = load_workbook(arquivo, data_only=False, keep_links=True)
        if "base" not in wb.sheetnames:
            raise GestorError("Excel", "Aba 'base' nao encontrada.")
        ws = wb["base"]
        _validar_estrutura_excel(ws)
        for loja, linha in LINHAS_EXCEL.items():
            ws[f"C{linha}"] = float(Decimal(str(valores[loja])))
        wb.save(temporario)
        wb.close()
        wb = None
        validar_excel_metas(temporario, valores)
        os.replace(temporario, arquivo)
        aguardar_excel_liberado(arquivo)
        gravados = validar_excel_metas(arquivo, valores)
    except Exception:
        if wb is not None:
            wb.close()
        temporario.unlink(missing_ok=True)
        raise
    if log:
        log_gestor(log, f"Backup do Excel: {backup}")
        log_gestor(log, "Excel salvo e validado: C19:C26.")
    return gravados, backup


def fechar_report_preview(preview, log):
    user32 = ctypes.windll.user32
    if _janela_existe(preview["hwnd"]):
        user32.PostMessageW(preview["hwnd"], 0x0010, 0, 0)
        try:
            _aguardar_fechar(preview["hwnd"], 15)
        except GestorError:
            raise GestorError(
                "Report Preview", "Nao foi possivel fechar apenas o Report Preview."
            )
    log_gestor(log, "Report Preview fechado com seguranca.")


def salvar_diagnostico_gestor(config, etapa, log, hwnd=None):
    config.pasta_erros.mkdir(parents=True, exist_ok=True)
    nome = re.sub(r"[^a-zA-Z0-9_-]+", "_", etapa).strip("_") or "gestor"
    caminho = config.pasta_erros / f"gestor_{nome}_{datetime.now(timezone.utc).strftime('%Y-%m-%d_%H-%M-%S')}.png"
    try:
        if hwnd and _janela_existe(hwnd):
            janela = next(
                (item for item in _enumerar_janelas() if item["hwnd"] == hwnd),
                None,
            )
            if janela:
                log_gestor(
                    log,
                    "Estado da janela no erro: "
                    f"titulo={janela['titulo']!r}; hwnd={hwnd}; "
                    f"classe={janela['classe']!r}; retangulo={janela['retangulo']}.",
                )
            _ativar(hwnd)
            time.sleep(0.5)
            _capturar_janela(hwnd, caminho)
        else:
            ImageGrab.grab(all_screens=True).save(caminho)
        log_gestor(log, f"Screenshot de diagnostico: {caminho}")
        return caminho
    except Exception as erro:
        log_gestor(log, f"Falha ao salvar screenshot de diagnostico: {erro}")
        return None


def executar_fase_gestor(config, log):
    inicio = time.monotonic()
    preview = None
    ultima_janela = None
    config.validar()
    config.pasta_logs.mkdir(parents=True, exist_ok=True)
    config.pasta_erros.mkdir(parents=True, exist_ok=True)
    config.pasta_backup.mkdir(parents=True, exist_ok=True)
    log_gestor(log, f"Inicio da Fase 3: {datetime.now(timezone.utc).strftime('%d/%m/%Y %H:%M:%S')}")
    try:
        menu = conectar_remoteapp(config, log)
        ultima_janela = menu
        gestor = abrir_lojas_hoje(config, log, menu=menu)
        ultima_janela = gestor
        local = autenticar_comercial(gestor, config, log)
        ultima_janela = local
        impressora = selecionar_local_99(local, config, log)
        ultima_janela = impressora
        modulo = selecionar_impressora_windows(impressora, config, log)
        ultima_janela = modulo
        liberacao = abrir_metas_vendas_local(modulo, config, log)
        ultima_janela = liberacao
        metas = autenticar_liberacao(liberacao, config, log)
        ultima_janela = metas
        preview, data_relatorio = gerar_relatorio_metas(metas, config, log)
        ultima_janela = preview
        evidencia = capturar_report_preview(preview, config, log)
        valores, total, _linhas = extrair_relatorio_com_tentativas(
            preview, evidencia, config, log
        )
        soma = validar_total_relatorio(valores, total)
        for loja in LOJAS:
            log_gestor(log, f"{loja}={formatar_brasileiro(valores[loja])}")
        log_gestor(log, f"Soma das lojas={formatar_brasileiro(soma)}")
        log_gestor(log, f"TOTAL GERAL={formatar_brasileiro(total)}")
        log_gestor(log, "Validacao do total=OK")
        atualizar_excel_metas(config.arquivo_excel, valores, config.pasta_backup, log)
        fechar_report_preview(preview, log)
        preview = None
        duracao = time.monotonic() - inicio
        log_gestor(log, f"FASE 3 CONCLUIDA COM SUCESSO em {duracao:.1f}s.")
        return ResultadoGestor(
            data_relatorio=data_relatorio,
            valores=valores,
            total_relatorio=total,
            soma_lojas=soma,
            arquivo_excel=config.arquivo_excel,
            screenshot=evidencia,
            paginas=1,
            duracao_segundos=duracao,
        )
    except Exception as erro:
        etapa = erro.etapa if isinstance(erro, GestorError) else "fase_gestor"
        hwnd = getattr(erro, "hwnd", None)
        if not hwnd and preview:
            hwnd = preview["hwnd"]
        if not hwnd and ultima_janela:
            hwnd = ultima_janela["hwnd"]
        salvar_diagnostico_gestor(config, etapa, log, hwnd=hwnd)
        if preview and _janela_existe(preview["hwnd"]):
            try:
                fechar_report_preview(preview, log)
            except Exception as erro_fechar:
                log_gestor(log, f"Falha ao fechar Report Preview apos erro: {erro_fechar}")
        if isinstance(erro, GestorError):
            raise
        raise GestorError(etapa, str(erro)) from erro
    log_gestor(log, "Fase 3 finalizada.")
    log_gestor(log, f"Fim da Fase 3: {datetime.now(timezone.utc).strftime('%d/%m/%Y %H:%M:%S')}")
    log_gestor(log, f"fechar_janela_gestor({gestor['hwnd']})")
    _fechar_janela_gestor(gestor["hwnd"])
    if _janela_existe(gestor["hwnd"]):
        log_gestor(log, "Falha ao fechar a janela do Gestor de Vendas.")
    else:
        log_gestor(log, "Janela do Gestor de Vendas fechada com sucesso.")
        try:
            _aguardar_fechar(gestor["hwnd"], 15)
        except GestorError:
            log_gestor(log, "Falha ao aguardar o fechamento da janela do Gestor de Vendas.")
            def _forcar_fechar(hwnd):
                user32 = ctypes.windll.user32
                user32.PostMessageW(hwnd, 0x0010, 0, 0)
            _forcar_fechar(gestor["hwnd"])
            def _aguardar_fechar_forcado(hwnd, timeout=15):
                fim = time.monotonic() + timeout
                while time.monotonic() < fim:
                    if not _janela_existe(hwnd):
                        return
                    time.sleep(0.5)
                raise GestorError("fechar_janela_gestor", "Falha ao forcar o fechamento da janela do Gestor de Vendas.")

