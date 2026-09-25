"""Normalizacao restrita a DATA CPC, sem deduplicar ou completar o historico."""

import argparse
from collections import Counter
from datetime import date, datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import shutil
import tempfile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def coluna_data_cpc(ws):
    colunas = [
        cell.column for cell in ws[1]
        if " ".join(str(cell.value or "").upper().split()) == "DATA CPC"
    ]
    if len(colunas) != 1:
        raise ValueError("Receita: esperado um unico cabecalho DATA CPC.")
    return colunas[0]


def converter_data_cpc(valor):
    if valor is None or isinstance(valor, str) and not valor.strip():
        return None
    if isinstance(valor, datetime):
        return valor
    if isinstance(valor, date):
        return datetime.combine(valor, datetime.min.time())
    if isinstance(valor, str):
        for formato in ("%d/%m/%Y", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
            try:
                return datetime.strptime(valor.strip(), formato)
            except ValueError:
                pass
    raise ValueError("DATA CPC invalida; esperado uma data real ou texto DD/MM/AAAA.")


def normalizar_datas_cpc(ws):
    coluna = coluna_data_cpc(ws)
    pendentes = []
    meses = Counter()
    for linha in ws.iter_rows(min_row=2, min_col=coluna, max_col=coluna):
        cell = linha[0]
        try:
            valor = converter_data_cpc(cell.value)
        except ValueError as erro:
            raise ValueError(f"Receita: DATA CPC invalida em {cell.coordinate}.") from erro
        if valor is not None:
            meses[valor.strftime("%Y-%m")] += 1
            pendentes.append((cell, valor))
    convertidas = sum(not isinstance(cell.value, (date, datetime)) for cell, _ in pendentes)
    # Valida toda a coluna antes da primeira alteracao.
    for cell, valor in pendentes:
        cell.value = valor
        cell.number_format = "dd/mm/yyyy"
    return {"coluna": get_column_letter(coluna), "convertidas": convertidas, "meses": dict(sorted(meses.items()))}


def preparar_linhas_receita(linhas, coluna):
    resultado = []
    for numero, linha in enumerate(linhas, start=2):
        if all(valor is None or str(valor).strip() == "" for valor in linha):
            continue
        if len(linha) < coluna:
            raise ValueError(f"Receita: exportacao sem DATA CPC na linha {numero}.")
        valores = list(linha)
        try:
            valores[coluna - 1] = converter_data_cpc(valores[coluna - 1])
        except ValueError as erro:
            raise ValueError(f"Receita: DATA CPC invalida na exportacao, linha {numero}.") from erro
        resultado.append(valores)
    if not resultado:
        raise ValueError("Receita: exportacao sem registros validos.")
    return resultado


def _assinatura_preservada(wb, aba, coluna, conferir_datas=None):
    assinatura = hashlib.sha256()
    for ws in wb:
        assinatura.update(ws.title.encode("utf-8"))
        for linha in ws:
            for cell in linha:
                if cell.value is None:
                    continue
                if ws.title == aba and cell.column == coluna and cell.row > 1:
                    if conferir_datas is not None:
                        if not isinstance(cell.value, datetime) or cell.number_format != "dd/mm/yyyy":
                            raise ValueError(f"Data nao normalizada: {cell.coordinate}.")
                        conferir_datas[cell.value.strftime("%Y-%m")] += 1
                    continue
                assinatura.update(repr((cell.coordinate, cell.value, cell.number_format)).encode("utf-8"))
    return assinatura.hexdigest()


def _hash_arquivo(arquivo):
    with Path(arquivo).open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def corrigir_arquivo(arquivo, pasta_backup):
    arquivo = Path(arquivo)
    # Nao salvar por cima de uma planilha aberta pelo usuario.
    if arquivo.with_name("~$" + arquivo.name).exists():
        raise PermissionError("Feche a planilha TESTE no Excel antes da correcao.")
    assinatura_original = _hash_arquivo(arquivo)
    print("[RECEITA] Lendo a planilha; original preservado ate a conferencia final.", flush=True)
    wb = load_workbook(arquivo, keep_links=True)
    temporario = None
    try:
        ws = wb.active
        if ws is None:
            raise ValueError("A planilha nao possui abas para processar.")
        coluna = coluna_data_cpc(ws)
        aba = ws.title
        assinatura = _assinatura_preservada(wb, aba, coluna)
        relatorio = normalizar_datas_cpc(ws)
        pasta_backup = Path(pasta_backup)
        pasta_backup.mkdir(parents=True, exist_ok=True)
        backup = pasta_backup / f"{arquivo.stem}_antes_datas_cpc_{datetime.now(timezone.utc):%Y%m%d_%H%M%S_%f}.xlsx"
        shutil.copy2(arquivo, backup)
        if _hash_arquivo(backup) != assinatura_original:
            raise RuntimeError("A planilha mudou durante a leitura; original preservado.")
        print(f"[RECEITA] Backup criado. Convertendo {relatorio['convertidas']} datas-texto.", flush=True)
        with tempfile.NamedTemporaryFile(dir=arquivo.parent, suffix=".xlsx", prefix=".cpc_", delete=False) as tmp:
            temporario = Path(tmp.name)
        wb.save(temporario)
        wb.close()
        print("[RECEITA] Conferindo datas, contagens e demais valores/formulas.", flush=True)
        conferencia = load_workbook(temporario, read_only=True, keep_links=True)
        try:
            meses = Counter()
            if _assinatura_preservada(conferencia, aba, coluna, conferir_datas=meses) != assinatura:
                raise ValueError("Conferencia falhou: conteudo fora de DATA CPC mudou.")
            if dict(meses) != relatorio["meses"]:
                raise ValueError("Conferencia falhou: contagem mensal mudou.")
        finally:
            conferencia.close()
        if _hash_arquivo(arquivo) != assinatura_original:
            raise RuntimeError("A planilha mudou durante a correcao; original preservado.")
        os.replace(temporario, arquivo)
        return {**relatorio, "arquivo": str(arquivo), "backup": str(backup)}
    finally:
        wb.close()
        if temporario is not None:
            temporario.unlink(missing_ok=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("arquivo", type=Path)
    parser.add_argument("--backup", type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(corrigir_arquivo(args.arquivo, args.backup), ensure_ascii=False, indent=2))
