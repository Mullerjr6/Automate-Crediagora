from pathlib import Path
from openpyxl import load_workbook

p = Path.home() / "TJI PROMOTORA DE VENDAS EIRELI" / "Crediagora-doc - dados" / "tabela fat" / "fat_metas_venda_moda_lojas_hoje.xlsx"
wb = load_workbook(p, data_only=False, read_only=True, keep_links=False)
for ws in wb.worksheets:
    print("ABA", ws.title, ws.max_row, ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 15), values_only=True):
        print(row)
    print("---")
wb.close()
