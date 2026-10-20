"""Prepara e valida cópias isoladas antes de migrar a automação às planilhas oficiais."""

import argparse
import json
import shutil
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

PASTA_TABELAS = (
    Path.home()
    / "TJI PROMOTORA DE VENDAS EIRELI"
    / "Crediagora-doc - dados"
    / "tabela fat"
)
PASTA_INDICADORES = (
    Path.home()
    / "TJI PROMOTORA DE VENDAS EIRELI"
    / "Crediagora-doc - dados"
    / "Indicadores"
)
PASTA_TESTE = Path(__file__).resolve().parent.parent / "teste_migracao_oficiais"
ARQUIVOS = {
    "vendas": (PASTA_TABELAS / "fat_vendas.xlsx", "fat_vendas_Teste.xlsx", ("U", "V")),
    "receita": (
        PASTA_TABELAS / "fat_receita_gerada_CPC.xlsx",
        "fat_receita_gerada_CPC_Teste.xlsx",
        ("W", "X", "Y", "Z", "AA"),
    ),
    "retorno_ligacoes": (
        PASTA_TABELAS / "fat_retorno_ligacoes.xlsx",
        "fat_retorno_ligacoes - teste automacao.xlsx",
        ("J", "K", "L"),
    ),
    "metas_gestor": (
        PASTA_TABELAS / "fat_metas_venda_moda_lojas_hoje.xlsx",
        "fat_metas_venda_moda_lojas_hoje.xlsx",
        (),
    ),
    "indicadores": (
        PASTA_INDICADORES / "Indicadores de FPD1.xlsx",
        "Indicadores de FPD1.xlsx",
        (),
    ),
}


def copiar_para_teste():
    """Cria cópias de trabalho sem modificar os arquivos oficiais."""
    PASTA_TESTE.mkdir(parents=True, exist_ok=True)
    copiados = []
    for nome, (origem, destino_nome, _) in ARQUIVOS.items():
        if not origem.is_file():
            raise FileNotFoundError(f"Planilha oficial não encontrada ({nome}): {origem}")
        destino = PASTA_TESTE / destino_nome
        shutil.copy2(origem, destino)
        copiados.append(str(destino))
    return copiados


def validar_arquivo(nome, arquivo, colunas_formula):
    resultado = {"arquivo": str(arquivo), "status": "aprovado", "erros": []}
    if not arquivo.is_file() or arquivo.stat().st_size == 0:
        resultado["status"] = "reprovado"
        resultado["erros"].append("arquivo inexistente ou vazio")
        return resultado
    try:
        with zipfile.ZipFile(arquivo) as pacote:
            problema = pacote.testzip()
            if problema:
                resultado["erros"].append(f"XML interno inválido: {problema}")
            links_externos = [item for item in pacote.namelist() if item.startswith("xl/externalLinks/")]
            if links_externos:
                resultado["erros"].append("contém vínculos externos do Excel")
        wb = load_workbook(arquivo, data_only=False, keep_links=False, read_only=True)
        resultado["abas"] = wb.sheetnames
        resultado["formulas"] = {}
        for ws in wb.worksheets:
            if ws.max_row < 1 or ws.max_column < 1:
                resultado["erros"].append(f"aba vazia: {ws.title}")
            quantidade = 0
            for linha in ws.iter_rows():
                quantidade += sum(celula.data_type == "f" for celula in linha)
            resultado["formulas"][ws.title] = quantidade
        if colunas_formula:
            ws = wb.active
            total = sum(resultado["formulas"].values())
            if total == 0:
                resultado["erros"].append(
                    f"nenhuma fórmula encontrada; esperadas fórmulas nas colunas {', '.join(colunas_formula)}"
                )
        wb.close()
    except (OSError, ValueError, zipfile.BadZipFile) as erro:
        resultado["erros"].append(f"não abre com segurança: {erro}")
    if resultado["erros"]:
        resultado["status"] = "reprovado"
    return resultado


def executar_validacao():
    relatorio = {
        "gerado_em": datetime.now().astimezone().isoformat(timespec="seconds"),
        "pasta_teste": str(PASTA_TESTE),
        "itens": {},
    }
    for nome, (_, destino_nome, colunas_formula) in ARQUIVOS.items():
        relatorio["itens"][nome] = validar_arquivo(
            nome,
            PASTA_TESTE / destino_nome,
            colunas_formula,
        )
    relatorio["aprovado"] = all(
        item["status"] == "aprovado" for item in relatorio["itens"].values()
    )
    destino = PASTA_TESTE / "relatorio_validacao.json"
    destino.write_text(json.dumps(relatorio, ensure_ascii=False, indent=2), encoding="utf-8")
    return relatorio, destino


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--preparar", action="store_true", help="Copia as oficiais para o ambiente isolado.")
    args = parser.parse_args()
    if args.preparar:
        for arquivo in copiar_para_teste():
            print(f"[COPIADO] {arquivo}")
    relatorio, destino = executar_validacao()
    print(f"[{'APROVADO' if relatorio['aprovado'] else 'REPROVADO'}] Relatório: {destino}")
    return 0 if relatorio["aprovado"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
