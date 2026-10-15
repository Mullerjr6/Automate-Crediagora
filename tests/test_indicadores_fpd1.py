import csv
import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.formula.translate import Translator

RAIZ = Path(__file__).resolve().parents[1]
MODULO = RAIZ / "CREDIAGORA" / "Script" / "indicadores_fpd1.py"


def carregar_modulo():
    spec = importlib.util.spec_from_file_location("indicadores_fpd1_testes", MODULO)
    modulo = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = modulo
    spec.loader.exec_module(modulo)
    return modulo


def criar_xlsx(caminho, total_linhas=5):
    wb = Workbook()
    ws = wb.active
    ws.title = "Tabela Conta"
    cabecalhos = [f"Campo{i}" for i in range(1, 51)]
    cabecalhos[0], cabecalhos[1], cabecalhos[12] = "DataContabil", "Proposta", "Cpf"
    ws.append(cabecalhos)
    for linha in range(2, total_linhas + 1):
        for coluna in range(1, 51):
            ws.cell(linha, coluna).value = f"antigo-{linha}-{coluna}"
        ws.cell(linha, 51).value = f"=A{linha}"
        ws.cell(linha, 52).value = f"=$A$1&A{linha}"
        ws.cell(linha, 53).value = 7
    wb.save(caminho)
    wb.close()
    return cabecalhos


def criar_csv(caminho, cabecalhos, registros):
    with caminho.open("w", encoding="utf-8-sig", newline="") as arquivo:
        escritor = csv.writer(arquivo, delimiter=";")
        escritor.writerow(cabecalhos + [""])
        for registro in registros:
            escritor.writerow(registro + [""])


class IndicadoresFpd1Tests(unittest.TestCase):
    def test_traducao_de_formulas_mantem_referencias_complexas(self):
        modulo = carregar_modulo()
        formulas = (
            '=TEXT(AO2,"AAAMM")',
            '=IF(OR(MONTH(AS2)-MONTH(AO2)>1,AS2=""),1,0)',
            '=FV($CM$2,(AO2-A2)/30,,-SUM(AI2+AK2+BQ2))',
            '=IFERROR(VLOOKUP(M2,[1]Dados!$L:$CH,75,FALSE),"")',
            '=IFERROR(VLOOKUP(B2,\'[3]exportacao-161-20240130135023\'!$J:$Y,16,0)," ")',
        )
        for formula in formulas:
            with self.subTest(formula=formula):
                esperado = Translator(formula, origin="AY2").translate_formula("AY19826")
                self.assertEqual(modulo._traduzir(formula, "AY2", "AY19826"), esperado)

    def test_reparo_reconstroi_somente_formulas_e_cria_backup(self):
        modulo = carregar_modulo()
        with tempfile.TemporaryDirectory() as pasta:
            pasta = Path(pasta)
            xlsx, backups = pasta / "indicadores.xlsx", pasta / "backups"
            criar_xlsx(xlsx)
            wb = load_workbook(xlsx)
            ws = wb["Tabela Conta"]
            ws["AY2"] = '=IFERROR(VLOOKUP(M2,[1]Dados!$L:$CH,75,FALSE),"")'
            ws["AZ2"] = '=IFERROR(VLOOKUP(B2,\'[3]exportacao-161-20240130135023\'!$J:$Y,16,0)," ")'
            ws["AY3"] = "=FORMULA_BUGADA()"
            ws["AZ3"] = "=FORMULA_BUGADA()"
            ws["CM3"] = "=FORMULA_LEGADA()"
            dado_original = ws["A3"].value
            wb.save(xlsx)
            wb.close()

            arquivo, backup, ultima, formulas = modulo.reparar_formulas_indicadores_fpd1(
                xlsx, backups, lambda _mensagem: None
            )

            self.assertEqual(arquivo, xlsx)
            self.assertTrue(backup.exists())
            self.assertEqual((ultima, formulas), (5, 2))
            wb = load_workbook(xlsx, data_only=False, keep_links=True)
            ws = wb["Tabela Conta"]
            self.assertEqual(ws["A3"].value, dado_original)
            self.assertEqual(
                ws["AY3"].value,
                Translator(ws["AY2"].value, origin="AY2").translate_formula("AY3"),
            )
            self.assertEqual(
                ws["AZ5"].value,
                Translator(ws["AZ2"].value, origin="AZ2").translate_formula("AZ5"),
            )
            self.assertIsNone(ws["CM3"].value)
            wb.close()

    def test_conferencia_arquivo_detecta_xlsx_corrompido(self):
        modulo = carregar_modulo()
        with tempfile.TemporaryDirectory() as pasta:
            arquivo = Path(pasta) / "corrompido.xlsx"
            arquivo.write_bytes(b"nao e um xlsx")
            with self.assertRaisesRegex(modulo.AtualizacaoIndicadoresError, "corrompido"):
                modulo._validar_arquivo_salvo(arquivo)

    def test_atualiza_formulas_textos_e_remove_sobras(self):
        modulo = carregar_modulo()
        with tempfile.TemporaryDirectory() as pasta:
            pasta = Path(pasta)
            xlsx, csv_atual = pasta / "indicadores.xlsx", pasta / "contratos.csv"
            cabecalhos = criar_xlsx(xlsx)
            wb = load_workbook(xlsx)
            ws = wb["Tabela Conta"]
            for linha in ws.iter_rows(min_row=4, max_row=5, max_col=50):
                for celula in linha:
                    celula.value = None
            ws["A6"] = "   "
            ws["AY7"] = "=A7"
            wb.save(xlsx)
            wb.close()
            registros = []
            for numero in (1, 2):
                linha = [str(numero)] * 50
                linha[0] = f"0{numero}/01/2024"
                linha[1] = f"000{numero}"
                linha[12] = f"0000000000{numero}"
                registros.append(linha)
            criar_csv(csv_atual, cabecalhos, registros)
            with patch.object(modulo.time, "sleep", return_value=None):
                resultado = modulo.atualizar_indicadores_fpd1(
                    csv_atual, xlsx, lambda _mensagem: None
                )
            self.assertEqual((resultado.registros_importados, resultado.ultima_linha), (2, 3))
            wb = load_workbook(xlsx, data_only=False)
            ws = wb["Tabela Conta"]
            self.assertEqual(ws["B2"].value, "0001")
            self.assertEqual(ws["M2"].value, "00000000001")
            self.assertEqual(ws["AY3"].value, "=A3")
            self.assertEqual(ws["AZ3"].value, "=$A$1&A3")
            self.assertEqual(ws["BA3"].value, 7)
            self.assertIsNone(ws["A4"].value)
            self.assertIsNone(ws["AY4"].value)
            wb.close()

    def test_csv_menor_bloqueia_sem_alterar_xlsx(self):
        modulo = carregar_modulo()
        with tempfile.TemporaryDirectory() as pasta:
            pasta = Path(pasta)
            xlsx, csv_atual = pasta / "indicadores.xlsx", pasta / "contratos.csv"
            cabecalhos = criar_xlsx(xlsx)
            criar_csv(csv_atual, cabecalhos, [["1"] * 50] * 3 + [[""] * 50] * 3)
            original = xlsx.read_bytes()
            mensagens = []
            with patch.object(modulo.time, "sleep", return_value=None), self.assertRaisesRegex(
                modulo.AtualizacaoIndicadoresError,
                "CSV possui 3 registros, menos que os 4",
            ):
                modulo.atualizar_indicadores_fpd1(csv_atual, xlsx, mensagens.append)
            self.assertEqual(xlsx.read_bytes(), original)
            self.assertTrue(csv_atual.exists())
            self.assertTrue(any("Atualizacao bloqueada" in mensagem for mensagem in mensagens))
            self.assertEqual(sorted(p.name for p in pasta.iterdir()), ["contratos.csv", "indicadores.xlsx"])

    def test_csv_igual_ou_maior_permite_atualizacao(self):
        modulo = carregar_modulo()
        for quantidade in (4, 5):
            with self.subTest(quantidade=quantidade), tempfile.TemporaryDirectory() as pasta:
                pasta = Path(pasta)
                xlsx, csv_atual = pasta / "indicadores.xlsx", pasta / "contratos.csv"
                cabecalhos = criar_xlsx(xlsx)
                registro = ["1"] * 50
                registro[0] = "01/01/2024"
                criar_csv(csv_atual, cabecalhos, [registro] * quantidade)
                with patch.object(modulo.time, "sleep", return_value=None):
                    resultado = modulo.atualizar_indicadores_fpd1(
                        csv_atual, xlsx, lambda _mensagem: None
                    )
                self.assertEqual(resultado.registros_importados, quantidade)

    def test_csv_invalido_nao_altera_original(self):
        modulo = carregar_modulo()
        with tempfile.TemporaryDirectory() as pasta:
            pasta = Path(pasta)
            xlsx, csv_atual = pasta / "indicadores.xlsx", pasta / "contratos.csv"
            cabecalhos = criar_xlsx(xlsx)
            criar_csv(csv_atual, cabecalhos, [["1"] * 50])
            with csv_atual.open("a", encoding="utf-8") as arquivo:
                arquivo.write(";EXTRA")
            original = xlsx.read_bytes()
            with patch.object(modulo.time, "sleep", return_value=None), self.assertRaises(
                modulo.AtualizacaoIndicadoresError
            ):
                modulo.atualizar_indicadores_fpd1(csv_atual, xlsx, lambda _mensagem: None)
            self.assertEqual(xlsx.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()
