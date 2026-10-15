import sys
import os
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "CREDIAGORA" / "Script"))
import receita_cpc as receita


class ReceitaTests(unittest.TestCase):
    def base(self):
        wb = Workbook()
        ws = wb.active
        ws["K1"] = "DATA CPC"
        ws["L1"] = "CPB"
        ws["K2"] = "01/05/2026"
        ws["K3"] = datetime(2026, 9, 25)
        ws["L2"] = "26/09/2026"
        ws["W2"] = "=WEEKDAY(K2,2)"
        return wb

    def test_converte_datas_sem_inverter_mes_e_dia(self):
        wb = self.base()
        resultado = receita.normalizar_datas_cpc(wb.active)
        self.assertEqual(resultado, {"coluna": "K", "convertidas": 1, "meses": {"2026-05": 1, "2026-09": 1}})
        self.assertEqual(wb.active["K2"].value, datetime(2026, 5, 1))
        self.assertEqual(wb.active["K2"].number_format, "dd/mm/yyyy")
        self.assertEqual(wb.active["L2"].value, "26/09/2026")
        self.assertEqual(wb.active["W2"].value, "=WEEKDAY(K2,2)")

    def test_rejeita_data_invalida_sem_alterar_coluna(self):
        wb = self.base()
        wb.active["K4"] = "31/02/2026"
        with self.assertRaisesRegex(ValueError, "K4"):
            receita.normalizar_datas_cpc(wb.active)
        self.assertEqual(wb.active["K2"].value, "01/05/2026")

    def test_cabecalho_duplicado_ou_ausente_interrompe(self):
        wb = self.base()
        wb.active["L1"] = "DATA CPC"
        with self.assertRaises(ValueError):
            receita.coluna_data_cpc(wb.active)
        wb.active["K1"] = wb.active["L1"] = "outro"
        with self.assertRaises(ValueError):
            receita.coluna_data_cpc(wb.active)

    def test_prepara_csv_preservando_demais_campos(self):
        linha = ["id"] * 10 + ["02/09/2026", "25/05/2026"]
        convertido = receita.preparar_linhas_receita([linha], 11)[0]
        self.assertEqual(convertido[10], datetime(2026, 9, 2))
        self.assertEqual(convertido[11], linha[11])
        self.assertEqual(linha[10], "02/09/2026")

    def test_reparo_backup_preserva_formulas_e_contagem(self):
        with tempfile.TemporaryDirectory() as pasta:
            arquivo = Path(pasta) / "teste.xlsx"
            wb = self.base()
            wb.create_sheet("outra")["A1"] = "preservar"
            wb.save(arquivo)
            wb.close()
            original = arquivo.read_bytes()
            resultado = receita.corrigir_arquivo(arquivo, Path(pasta) / "backups")
            self.assertEqual(Path(resultado["backup"]).read_bytes(), original)
            final = load_workbook(arquivo)
            self.assertEqual(final.active["K2"].value, datetime(2026, 5, 1))
            self.assertEqual(final.active["W2"].value, "=WEEKDAY(K2,2)")
            self.assertEqual(final["outra"]["A1"].value, "preservar")
            final.close()

    def test_falha_na_conferencia_preserva_original(self):
        with tempfile.TemporaryDirectory() as pasta:
            arquivo = Path(pasta) / "teste.xlsx"
            self.base().save(arquivo)
            original = arquivo.read_bytes()
            with patch.object(receita, "_assinatura_preservada", side_effect=["antes", "diferente"]):
                with self.assertRaisesRegex(ValueError, "conteudo"):
                    receita.corrigir_arquivo(arquivo, Path(pasta) / "backups")
            self.assertEqual(arquivo.read_bytes(), original)
            self.assertFalse(list(Path(pasta).glob(".cpc_*.xlsx")))

    def test_quantidade_mes_vigente_protege_vendas_e_receita(self):
        import start

        hoje = datetime.today()
        for coluna, dados, formulas, total in (
            (11, "V", ["W"], "Y"),
            (12, "T", ["U", "V"], "V"),
        ):
            for quantidade in (1, 2, 3):
                with self.subTest(coluna=coluna, quantidade=quantidade), tempfile.TemporaryDirectory() as pasta:
                    destino = Path(pasta) / "teste.xlsx"
                    origem = Path(pasta) / "exportacao.csv"
                    origem.touch()
                    wb = Workbook()
                    ws = wb.active
                    ws.cell(1, coluna, "DATA CPC")
                    ws["A2"] = "historico"
                    ws.cell(2, coluna, datetime(hoje.year - 1, hoje.month, 1))
                    for numero in (3, 4):
                        ws.cell(numero, 1, "antigo")
                        ws.cell(numero, coluna, hoje.strftime("%d/%m/%Y"))
                    ws[f"{formulas[0]}2"] = "=A2"
                    ws[f"{formulas[0]}8"] = "=A8"
                    wb.save(destino)
                    wb.close()
                    original = destino.read_bytes()
                    linha = [None] * start.coluna_para_numero(dados)
                    linha[0], linha[coluna - 1] = "novo", hoje.strftime("%d/%m/%Y")
                    linhas = [linha] * quantidade + [[" "] * len(linha)]
                    with patch.object(start, "ler_linhas_exportacao", return_value=linhas), \
                         patch.object(start, "log"), \
                         patch.object(start, "limpar_e_reconstruir_base", wraps=start.limpar_e_reconstruir_base) as limpar:
                        if quantidade < 2:
                            with self.assertRaisesRegex(ValueError, "Atualizacao bloqueada"):
                                start.copiar_dados_excel_origem_para_destino(
                                    origem, destino, colunas_formula=formulas,
                                    ultima_coluna_dados=dados, ultima_coluna_total=total,
                                    normalizar_cpc=True,
                                )
                            limpar.assert_not_called()
                            self.assertEqual(destino.read_bytes(), original)
                            self.assertTrue(origem.exists())
                        else:
                            start.copiar_dados_excel_origem_para_destino(
                                origem, destino, colunas_formula=formulas,
                                ultima_coluna_dados=dados, ultima_coluna_total=total,
                                normalizar_cpc=True,
                            )
                            limpar.assert_called_once()
                            final = load_workbook(destino)
                            self.assertEqual(final.active["A2"].value, "historico")
                            atuais = sum(
                                start.linha_eh_mes_vigente(valores[0])
                                for valores in final.active.iter_rows(
                                    min_row=2, min_col=coluna, max_col=coluna, values_only=True
                                )
                            )
                            self.assertEqual(atuais, quantidade)
                            final.close()

    def test_indice_cpc_invalido_bloqueia_antes_de_preparar_dados(self):
        import start

        for indice in (None, 0, -1, "K"):
            with self.subTest(indice=indice), tempfile.TemporaryDirectory() as pasta:
                destino = Path(pasta) / "teste.xlsx"
                origem = Path(pasta) / "exportacao.csv"
                origem.touch()
                wb = self.base()
                wb.save(destino)
                wb.close()
                original = destino.read_bytes()
                with patch.object(start, "ler_linhas_exportacao", return_value=[["dado"] * 22]), \
                     patch.object(start, "coluna_data_cpc", return_value=indice), \
                     patch.object(start, "preparar_linhas_receita") as preparar, \
                     patch.object(start, "log"):
                    with self.assertRaisesRegex(ValueError, "Coluna de comparação inválida"):
                        start.copiar_dados_excel_origem_para_destino(
                            origem, destino, normalizar_cpc=True,
                        )
                preparar.assert_not_called()
                self.assertEqual(destino.read_bytes(), original)

    def test_importacao_remove_mes_por_data_cpc_e_nao_por_cpb(self):
        import start

        with tempfile.TemporaryDirectory() as pasta:
            destino = Path(pasta) / "receita.xlsx"
            origem = Path(pasta) / "exportacao.csv"
            origem.touch()
            wb = self.base()
            wb.active["A2"] = "historico"
            wb.active["A3"] = "substituir"
            wb.active["L3"] = datetime(2026, 5, 1)
            wb.save(destino)
            wb.close()
            linha = [None] * 22
            linha[0], linha[10], linha[11] = "novo", "25/09/2026", "01/05/2026"
            with patch.object(start, "ler_linhas_exportacao", return_value=[linha]), \
                 patch.object(start, "log"), \
                 patch.object(start, "linha_eh_mes_vigente", side_effect=lambda valor: (
                     isinstance(valor, datetime) and valor.year == 2026 and valor.month == 9
                 )):
                for _ in range(2):
                    start.copiar_dados_excel_origem_para_destino(
                        origem, destino, colunas_formula=["W"],
                        ultima_coluna_dados="V", ultima_coluna_total="Y",
                        normalizar_cpc=True,
                    )
            final = load_workbook(destino)
            self.assertEqual(final.active["A2"].value, "historico")
            self.assertEqual(final.active["A3"].value, "novo")
            self.assertEqual(final.active["K3"].value, datetime(2026, 9, 25))
            self.assertEqual(final.active["L3"].value, "01/05/2026")
            self.assertEqual(final.active.max_row, 3)
            final.close()

    def test_mudanca_apenas_no_horario_nao_bloqueia_correcao(self):
        with tempfile.TemporaryDirectory() as pasta:
            arquivo = Path(pasta) / "teste.xlsx"
            self.base().save(arquivo)
            assinar = receita._assinatura_preservada

            def conferir(*args, **kwargs):
                if kwargs.get("conferir_datas") is not None:
                    estado = arquivo.stat()
                    os.utime(arquivo, (estado.st_atime, estado.st_mtime + 10))
                return assinar(*args, **kwargs)

            with patch.object(receita, "_assinatura_preservada", side_effect=conferir):
                resultado = receita.corrigir_arquivo(arquivo, Path(pasta) / "backups")
            self.assertEqual(resultado["convertidas"], 1)

    def test_vendas_data_cpc_em_l_preserva_cadastro_e_formulas(self):
        import start

        with tempfile.TemporaryDirectory() as pasta:
            destino = Path(pasta) / "vendas_teste.xlsx"
            origem = Path(pasta) / "vendas.csv"
            origem.touch()
            wb = Workbook()
            ws = wb.active
            ws["K1"], ws["L1"] = "Data Cadastro", "data CPC"
            ws["A2"], ws["A3"] = "historico", "antigo"
            ws["L2"], ws["L3"] = "01/05/2026", datetime(2026, 9, 1)
            ws["U2"], ws["V2"] = "=WEEKDAY(L2,2)", "=U2"
            wb.save(destino)
            wb.close()
            linha = [None] * 20
            linha[0], linha[10], linha[11] = "novo", "20/09/2026", "25/09/2026"
            with patch.object(start, "ler_linhas_exportacao", return_value=[linha]), \
                 patch.object(start, "log"), \
                 patch.object(start, "load_workbook", wraps=load_workbook) as abrir, \
                 patch.object(start, "linha_eh_mes_vigente", side_effect=lambda valor: (
                     isinstance(valor, datetime) and valor.year == 2026 and valor.month == 9
                 )):
                for _ in range(2):
                    start.copiar_dados_excel_origem_para_destino(origem, destino, normalizar_cpc=True)
                self.assertTrue(all(c.kwargs["keep_links"] for c in abrir.call_args_list))
            final = load_workbook(destino)
            self.assertEqual(final.active.max_row, 3)
            self.assertEqual(final.active["L2"].value, datetime(2026, 5, 1))
            self.assertEqual(final.active["L3"].value, datetime(2026, 9, 25))
            self.assertEqual(final.active["L3"].number_format, "dd/mm/yyyy")
            self.assertEqual(final.active["K3"].value, "20/09/2026")
            self.assertEqual(final.active["U3"].value, "=WEEKDAY(L3,2)")
            final.close()

    def test_fluxo_vendas_habilita_normalizacao_cpc(self):
        import start

        with patch.object(start, "baixar_exportacao", return_value=Path("vendas.csv")), \
             patch.object(start, "copiar_dados_excel_origem_para_destino") as copiar, \
             patch.object(start, "log"):
            start.processar_exportacao_vendas(object())
        self.assertTrue(copiar.call_args.kwargs["normalizar_cpc"])

    def test_fluxo_receita_aplica_formulas_ate_aa(self):
        import start

        with patch.object(start, "baixar_exportacao", return_value=Path("receita.csv")), \
             patch.object(start, "copiar_dados_excel_origem_para_destino") as copiar, \
             patch.object(start, "log"):
            start.processar_receita_gerada(object())

        self.assertEqual(
            copiar.call_args.kwargs["colunas_formula"], ["W", "X", "Y", "Z", "AA"]
        )
        self.assertEqual(copiar.call_args.kwargs["ultima_coluna_total"], "AA")

    def test_mudanca_real_durante_correcao_nao_e_sobrescrita(self):
        with tempfile.TemporaryDirectory() as pasta:
            arquivo = Path(pasta) / "teste.xlsx"
            self.base().save(arquivo)
            assinar = receita._assinatura_preservada

            def conferir(*args, **kwargs):
                if kwargs.get("conferir_datas") is not None:
                    with arquivo.open("ab") as stream:
                        stream.write(b"mudanca externa")
                return assinar(*args, **kwargs)

            with patch.object(receita, "_assinatura_preservada", side_effect=conferir):
                with self.assertRaisesRegex(RuntimeError, "mudou"):
                    receita.corrigir_arquivo(arquivo, Path(pasta) / "backups")
            self.assertTrue(arquivo.read_bytes().endswith(b"mudanca externa"))


if __name__ == "__main__":
    unittest.main()
