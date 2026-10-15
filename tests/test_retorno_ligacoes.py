import csv
import io
import sys
import tempfile
import unittest
import zipfile
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "CREDIAGORA" / "Script"))
import retorno_ligacoes as retorno


CABECALHO = ["CPF", "cliente", "Campanha", "estabelecimento", "Data_Contrato",
             "Status", "Usuário", "Telefone", "Celular"]


class RetornoLigacoesTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.pasta = Path(self.tmp.name)
        self.destino = self.pasta / "destino.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Base"
        ws.append(CABECALHO + ["Situação", "Considerar", "AnoMes", "Preservar"])
        for n in range(2, 6):
            ws.append([f"antigo{n}"] * 9 + [f'=IF(F{n}="","",validador!$A$2)',
                                           f'=J{n}&$M$1', f'=YEAR(E{n})', "fora"])
            ws.cell(n, 1).fill = PatternFill("solid", fgColor="ABCDEF")
        ws.auto_filter.ref = "A1:L5"
        wb.create_sheet("validador").append(["Status", "Regra"])
        wb["validador"].append(["A", "Sim"])
        wb["validador"].append(["Prospecção", "Prospecção", "Ligação"])
        wb.save(self.destino)
        wb.close()

    def criar_zip(self, linhas, nome="exportacao.zip", codificacao="utf-8-sig", delimitador=";"):
        saida = io.StringIO(newline="")
        escritor = csv.writer(saida, delimiter=delimitador)
        escritor.writerow(CABECALHO)
        escritor.writerows(linhas)
        caminho = self.pasta / nome
        with zipfile.ZipFile(caminho, "w") as arquivo:
            arquivo.writestr("pasta/contatos.csv", saida.getvalue().encode(codificacao))
        return caminho

    def registro(self, cpf="00123456789"):
        return [cpf, "João; teste", "Normal", "0001", "15/10/2026", "Ativo", "usuario", "0011", "0009"]

    def test_substituicao_menor_preserva_cabecalho_estilos_outras_abas_e_colunas(self):
        registros = [self.registro(), self.registro("=1+1")]
        arquivo = self.criar_zip(registros, codificacao="cp1252")
        retorno.atualizar_retorno_ligacoes(arquivo, self.destino, lambda _: None)
        wb = load_workbook(self.destino)
        self.addCleanup(wb.close)
        ws = wb["Base"]
        self.assertEqual([ws.cell(1, c).value for c in range(1, 10)], CABECALHO)
        self.assertEqual([ws.cell(2, c).value for c in range(1, 10)], registros[0])
        self.assertEqual(ws["A3"].value, "=1+1")
        self.assertEqual(ws["A3"].data_type, "s")
        self.assertEqual(ws["J3"].value, '=IFERROR(INDEX(validador!$B$2:$B$13,MATCH(F3,validador!$A$2:$A$13,0)),"")')
        self.assertEqual(ws["K3"].value, '=IFERROR(INDEX(validador!$C$2:$C$13,MATCH(J3,validador!$B$2:$B$13,0)),"")')
        self.assertEqual(ws["L3"].value, '=YEAR(E3)')
        self.assertTrue(all(c.value is None for linha in ws.iter_rows(min_row=4, max_row=5, max_col=12) for c in linha))
        self.assertEqual(ws["A5"].fill.fgColor.rgb, "00ABCDEF")
        self.assertEqual(ws["M5"].value, "fora")
        self.assertEqual(wb["validador"]["B2"].value, "Sim")
        self.assertEqual(ws.auto_filter.ref, "A1:L3")

    def test_base_maior_estende_formulas_e_estilo(self):
        arquivo = self.criar_zip([self.registro()] * 7, delimitador=",")
        retorno.atualizar_retorno_ligacoes(arquivo, self.destino, lambda _: None)
        wb = load_workbook(self.destino)
        self.addCleanup(wb.close)
        self.assertEqual(wb["Base"]["L8"].value, "=YEAR(E8)")
        self.assertEqual(wb["Base"]["A8"]._style, wb["Base"]["A2"]._style)

    def test_erro_de_salvamento_preserva_original(self):
        arquivo = self.criar_zip([self.registro()])
        antes = self.destino.read_bytes()
        with patch.object(retorno.os, "replace", side_effect=PermissionError("Excel bloqueado")):
            with self.assertRaisesRegex(retorno.RetornoLigacoesError, "ETAPA 7"):
                retorno.atualizar_retorno_ligacoes(arquivo, self.destino, lambda _: None)
        self.assertEqual(antes, self.destino.read_bytes())
        self.assertEqual(list(self.pasta.glob("*.tmp.xlsx")), [])

    def test_csv_vazio_ou_malformado_preserva_original(self):
        antes = self.destino.read_bytes()
        for linhas in ([], [["faltam", "colunas"]]):
            with self.subTest(linhas=linhas):
                arquivo = self.criar_zip(linhas)
                with self.assertRaises(retorno.RetornoLigacoesError):
                    retorno.atualizar_retorno_ligacoes(arquivo, self.destino, lambda _: None)
                self.assertEqual(antes, self.destino.read_bytes())

    def test_formula_ausente_preserva_original(self):
        wb = load_workbook(self.destino)
        for linha in range(2, 6):
            wb["Base"].cell(linha, 11).value = None
        wb.save(self.destino)
        wb.close()
        antes = self.destino.read_bytes()
        with self.assertRaisesRegex(retorno.RetornoLigacoesError, "Fórmula-modelo ausente"):
            retorno.atualizar_retorno_ligacoes(self.criar_zip([self.registro()]), self.destino, lambda _: None)
        self.assertEqual(antes, self.destino.read_bytes())

    def test_csv_real_alias_de_data_e_tipo_excel(self):
        wb = load_workbook(self.destino)
        wb["Base"]["E2"].number_format = "m/d/yy h:mm"
        wb.save(self.destino)
        wb.close()
        registro = self.registro()
        registro[1] = "Texto\ufffd original"
        registro[4] = "2026-10-15 10:25:30"
        cabecalho = CABECALHO.copy()
        cabecalho[4] = "datacontato"
        caminho = self.pasta / "real.zip"
        with zipfile.ZipFile(caminho, "w") as arquivo:
            arquivo.writestr("dados.csv", ";".join(cabecalho) + "\n" + ";".join(registro))
        retorno.atualizar_retorno_ligacoes(caminho, self.destino, lambda _: None)
        wb = load_workbook(self.destino)
        self.addCleanup(wb.close)
        self.assertEqual(wb["Base"]["E2"].value, datetime(2026, 10, 15, 10, 25, 30))
        self.assertEqual(wb["Base"]["E1"].value, "Data_Contrato")
        self.assertEqual(wb["Base"]["B2"].value, registro[1])
        self.assertEqual(wb["Base"]["A2"].value, registro[0])

    def test_status_com_caracteres_de_substituicao_e_normalizado(self):
        registro = self.registro()
        registro[5] = "Prospec\ufffd\ufffdo"
        retorno.atualizar_retorno_ligacoes(self.criar_zip([registro]), self.destino, lambda _: None)
        wb = load_workbook(self.destino, data_only=False)
        self.addCleanup(wb.close)
        esperado = "Prospec" + chr(0xE7) + chr(0xE3) + "o"
        self.assertEqual(wb["Base"]["F2"].value, esperado)
        self.assertEqual(wb["Base"]["J2"].value, '=IFERROR(INDEX(validador!$B$2:$B$13,MATCH(F2,validador!$A$2:$A$13,0)),"")')

    def test_data_invalida_nao_altera_excel(self):
        wb = load_workbook(self.destino)
        wb["Base"]["E2"].number_format = "dd/mm/yyyy"
        wb.save(self.destino)
        wb.close()
        antes = self.destino.read_bytes()
        registro = self.registro()
        registro[4] = "nao e data"
        with self.assertRaisesRegex(retorno.RetornoLigacoesError, "Data inválida"):
            retorno.atualizar_retorno_ligacoes(self.criar_zip([registro]), self.destino, lambda _: None)
        self.assertEqual(antes, self.destino.read_bytes())

    def test_abas_ambiguas_bloqueiam_atualizacao(self):
        wb = load_workbook(self.destino)
        wb.copy_worksheet(wb["Base"])
        wb.save(self.destino)
        wb.close()
        antes = self.destino.read_bytes()
        with self.assertRaisesRegex(retorno.RetornoLigacoesError, "aba única"):
            retorno.atualizar_retorno_ligacoes(self.criar_zip([self.registro()]), self.destino, lambda _: None)
        self.assertEqual(antes, self.destino.read_bytes())

    def test_zip_corrompido_ou_sem_csv(self):
        caminho = self.pasta / "invalido.zip"
        caminho.write_bytes(b"zip corrompido")
        with self.assertRaises(retorno.RetornoLigacoesError):
            retorno.ler_csv_zip(caminho)
        with zipfile.ZipFile(caminho, "w") as arquivo:
            arquivo.writestr("arquivo.txt", "sem csv")
        with self.assertRaisesRegex(retorno.RetornoLigacoesError, "encontrados 0"):
            retorno.ler_csv_zip(caminho)

    def test_download_ignora_zip_antigo_e_aguarda_fim_do_novo(self):
        antigo = self.criar_zip([self.registro()], nome="antigo.zip")
        anteriores = retorno.inventariar_downloads(self.pasta)
        novo = self.criar_zip([self.registro()], nome="novo.zip")
        parcial = self.pasta / "novo.zip.crdownload"
        parcial.write_bytes(b"baixando")
        with patch.object(retorno.time, "sleep", side_effect=lambda _: parcial.unlink(missing_ok=True)):
            escolhido = retorno.esperar_zip_novo(self.pasta, anteriores, 2)
        self.assertEqual(escolhido, novo)
        self.assertTrue(antigo.exists())

    def test_download_nao_aceita_somente_zip_antigo(self):
        self.criar_zip([self.registro()])
        anteriores = retorno.inventariar_downloads(self.pasta)
        with patch.object(retorno.time, "monotonic", side_effect=[0, 0, 3]), patch.object(retorno.time, "sleep"):
            with self.assertRaisesRegex(retorno.RetornoLigacoesError, "ZIP novo não concluído"):
                retorno.esperar_zip_novo(self.pasta, anteriores, 2)

    def test_download_com_uuid_preserva_zip_com_nome_original(self):
        antigo = self.criar_zip([self.registro()], nome="exportacao.zip")
        conteudo = antigo.read_bytes()
        anteriores = retorno.inventariar_downloads(self.pasta)
        novo = self.criar_zip([self.registro("00000000001")], nome="f308a13e-93a8-4e70-8cbf-2a2c653240a5")
        with patch.object(retorno.time, "sleep"):
            self.assertEqual(retorno.esperar_zip_novo(self.pasta, anteriores, 2), novo)
        self.assertEqual(antigo.read_bytes(), conteudo)


if __name__ == "__main__":
    unittest.main()
