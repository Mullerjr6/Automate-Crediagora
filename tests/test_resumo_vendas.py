import configparser
import sys
import tempfile
import unittest
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "CREDIAGORA" / "Script"))
import resumo_vendas as resumo


class ResumoVendasTests(unittest.TestCase):
    def test_soma_apenas_mes_vigente_com_precisao_decimal(self):
        for tipo, campo in (("VENDAS", "Valor Contrato"), ("RECEITA", "Comissão Prevista")):
            with self.subTest(tipo=tipo):
                resultado = resumo.calcular_resumo(
                    [["14/10/2026", "R$ 1.234,56"], ["01/10/2026", 0.1],
                     ["02/10/2026", "0,20"], ["30/09/2026", "999,00"],
                     ["03/10/2026", "-4,86"]],
                    ["data CPC", campo], tipo, datetime(2026, 10, 14),
                )
                self.assertEqual(resultado.total, Decimal("1230.00"))
                self.assertEqual(resultado.registros, 4)

    def test_valores_monetarios(self):
        for entrada, esperado in (("1.234", "1234"), ("1234.56", "1234.56"),
                                  ("1.234.567,89", "1234567.89"), (None, "0")):
            self.assertEqual(resumo._valor_monetario(entrada), Decimal(esperado))
        for entrada in ("abc", "NaN", "Infinity", True, "1,234.56"):
            with self.subTest(entrada=entrada), self.assertRaises((TypeError, ValueError)):
                resumo._valor_monetario(entrada)

    def test_cabecalho_ausente_ou_duplicado_bloqueia(self):
        for campos in (["DATA CPC"], ["DATA CPC", "Valor Contrato", "VALOR CONTRATO"]):
            with self.assertRaises(ValueError):
                resumo.calcular_resumo([], campos, "VENDAS", datetime(2026, 10, 14))

    def test_um_txt_por_dia_preserva_outra_base_e_substitui_total(self):
        with tempfile.TemporaryDirectory() as pasta:
            primeiro = resumo.ResumoMensal("VENDAS", "10/2026", 2, Decimal("1000.10"))
            receita = resumo.ResumoMensal("RECEITA", "10/2026", 2, Decimal("40.20"))
            caminho = resumo.salvar_resumo(pasta, primeiro, "vendas.xlsx", datetime(2026, 10, 14, 8, 30))
            resumo.salvar_resumo(pasta, receita, "receita.xlsx", datetime(2026, 10, 14, 9, 15))
            resumo.salvar_resumo(pasta, primeiro, "vendas.xlsx", datetime(2026, 10, 14, 10, 45))
            documento = configparser.ConfigParser(interpolation=None)
            documento.read(caminho, encoding="utf-8")
            self.assertEqual(caminho.name, "14-10-2026.txt")
            self.assertEqual(documento["VENDAS"]["Soma do valor dos contratos"], "R$ 1.000,10")
            self.assertEqual(documento["RECEITA"]["Soma da comissao prevista"], "R$ 40,20")
            self.assertIn("10:45:00", documento["VENDAS"]["Atualizado em"])
            self.assertIn("09:15:00", documento["RECEITA"]["Atualizado em"])
            resumo.salvar_resumo(pasta, primeiro, "vendas.xlsx", datetime(2026, 10, 15))
            self.assertEqual(len(list(Path(pasta).glob("*.txt"))), 2)

    def test_falha_ao_substituir_txt_preserva_arquivo_anterior(self):
        with tempfile.TemporaryDirectory() as pasta:
            item = resumo.ResumoMensal("VENDAS", "10/2026", 1, Decimal("20"))
            caminho = resumo.salvar_resumo(pasta, item, "teste.xlsx", datetime(2026, 10, 14))
            anterior = caminho.read_bytes()
            with patch.object(resumo.os, "replace", side_effect=PermissionError):
                with self.assertRaises(PermissionError):
                    resumo.salvar_resumo(pasta, item, "teste.xlsx", datetime(2026, 10, 14))
            self.assertEqual(caminho.read_bytes(), anterior)
            self.assertFalse(list(Path(pasta).glob("*.tmp")))

    def test_txt_somente_depois_de_salvar_planilha(self):
        import start

        for falhar in (True, False):
            with self.subTest(falhar=falhar), tempfile.TemporaryDirectory() as pasta:
                pasta = Path(pasta)
                origem, destino = pasta / "origem.csv", pasta / "destino.xlsx"
                origem.touch()
                wb = Workbook()
                ws = wb.active
                ws["L1"], ws["P1"], ws["U2"] = "DATA CPC", "Valor Contrato", "=A2"
                wb.save(destino)
                wb.close()
                linha = [None] * 20
                linha[0], linha[11], linha[15] = "contrato", datetime.now(), "1.234,56"
                salvar = Workbook.save

                def gravar(planilha, caminho):
                    if falhar:
                        raise PermissionError("Arquivo bloqueado")
                    self.assertFalse((pasta / "resumos").exists())
                    return salvar(planilha, caminho)

                with patch.object(start, "ler_linhas_exportacao", return_value=[linha]), \
                     patch.object(start, "PASTA_RESUMOS_VENDAS", pasta / "resumos"), \
                     patch.object(start, "log"), patch.object(Workbook, "save", gravar):
                    if falhar:
                        with self.assertRaises(PermissionError):
                            start.copiar_dados_excel_origem_para_destino(
                                origem, destino, normalizar_cpc=True, tipo_resumo="VENDAS",
                            )
                        self.assertFalse((pasta / "resumos").exists())
                    else:
                        start.copiar_dados_excel_origem_para_destino(
                            origem, destino, normalizar_cpc=True, tipo_resumo="VENDAS",
                        )
                        txt = next((pasta / "resumos").glob("*.txt"))
                        self.assertIn("R$ 1.234,56", txt.read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
