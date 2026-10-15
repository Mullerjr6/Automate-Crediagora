import importlib.util
import sys
import tempfile
import unittest
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import ANY, call, patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from PIL import Image, ImageDraw


RAIZ = Path(__file__).resolve().parents[1]
SCRIPT_DIR = RAIZ / "CREDIAGORA" / "Script"
GESTOR_PY = SCRIPT_DIR / "gestor.py"

if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))


def carregar_gestor():
    spec = importlib.util.spec_from_file_location("gestor_testes", GESTOR_PY)
    modulo = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = modulo
    spec.loader.exec_module(modulo)
    return modulo


def criar_planilha(caminho, rotulo_errado=False, rotulos_por_formula=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "base"
    ws["A18"] = "Local"
    ws["C18"] = "Venda dia"
    lojas = ["05-TUC", "13-GMI", "16-DMC", "17-RGO", "10-CMI", "11-CAT", "09-SPM", "01-SCR"]
    if rotulo_errado:
        lojas[4] = "10-OUTRA"
    for linha, loja in enumerate(lojas, start=19):
        if rotulos_por_formula:
            origem = linha - 17
            ws[f"A{origem}"] = loja
            ws[f"A{linha}"] = f"=A{origem}"
        else:
            ws[f"A{linha}"] = loja
        ws[f"C{linha}"] = 1
        ws[f"D{linha}"] = f"=C{linha}*2"
        ws[f"C{linha}"].fill = PatternFill("solid", fgColor="00FF00")
    ws["Z50"] = "nao alterar"
    outra = wb.create_sheet("outra")
    outra["A1"] = "preservado"
    wb.save(caminho)
    wb.close()


class GestorTests(unittest.TestCase):
    def setUp(self):
        self.modulo = carregar_gestor()
        self.valores = {
            "TUC": Decimal("2538.00"),
            "GMI": Decimal("3188.59"),
            "DMC": Decimal("5368.74"),
            "RGO": Decimal("4507.15"),
            "CMI": Decimal("2651.70"),
            "CAT": Decimal("7779.06"),
            "SPM": Decimal("486.80"),
            "SCR": Decimal("3346.49"),
        }

    def test_numero_brasileiro(self):
        self.assertEqual(
            self.modulo.numero_brasileiro("29.866,53"),
            Decimal("29866.53"),
        )
        self.assertEqual(
            self.modulo.formatar_brasileiro(Decimal("29866.53")),
            "29.866,53",
        )

    def test_capslock_e_desativado_sem_enviar_credenciais(self):
        teclado = unittest.mock.MagicMock()
        win = unittest.mock.MagicMock()
        win.user32.GetKeyState.side_effect = [1, 0]
        with patch.dict(sys.modules, {"pyautogui": teclado}), \
             patch.object(self.modulo.ctypes, "windll", win, create=True), \
             patch.object(self.modulo, "_ativar"), \
             patch.object(self.modulo.time, "sleep"):
            self.modulo._preparar_teclado_credenciais(123, lambda _: None)
        teclado.press.assert_called_once_with("capslock")
        teclado.write.assert_not_called()

    def test_capslock_desligado_nao_e_alternado(self):
        teclado = unittest.mock.MagicMock()
        win = unittest.mock.MagicMock()
        win.user32.GetKeyState.return_value = 0
        with patch.dict(sys.modules, {"pyautogui": teclado}), \
             patch.object(self.modulo.ctypes, "windll", win, create=True), \
             patch.object(self.modulo, "_ativar"):
            self.modulo._preparar_teclado_credenciais(123, lambda _: None)
        teclado.press.assert_not_called()

    def test_capslock_persistente_interrompe_login(self):
        teclado = unittest.mock.MagicMock()
        win = unittest.mock.MagicMock()
        win.user32.GetKeyState.return_value = 1
        with patch.dict(sys.modules, {"pyautogui": teclado}), \
             patch.object(self.modulo.ctypes, "windll", win, create=True), \
             patch.object(self.modulo, "_ativar"), \
             patch.object(self.modulo.time, "sleep"):
            with self.assertRaisesRegex(self.modulo.GestorError, "credenciais nao enviadas"):
                self.modulo._preparar_teclado_credenciais(123, lambda _: None)
        teclado.press.assert_called_once_with("capslock")

    def test_senha_invalida_interrompe_antes_de_selecionar_empresa(self):
        with patch.object(self.modulo, "_janela_existe", return_value=True), \
             patch.object(self.modulo, "_retangulo_visivel", return_value=(0, 0, 900, 600)), \
             patch.object(self.modulo.ImageGrab, "grab"), \
             patch.object(self.modulo, "_ativar") as ativar, \
             patch.object(self.modulo, "_executar_ocr", return_value=[
                 {"text": "SENHA INVÁLIDA OU NÃO LOCALIZADA... Lojas Hoje"}
             ]):
            with self.assertRaisesRegex(self.modulo.GestorError, "sem nova tentativa"):
                self.modulo._aguardar_empresas_acesso(123, 10)
        ativar.assert_not_called()

    def test_login_aguarda_empresa_em_vez_de_sleep_fixo(self):
        with patch.object(self.modulo, "_janela_existe", return_value=True), \
             patch.object(self.modulo, "_retangulo_visivel", return_value=(0, 0, 900, 600)), \
             patch.object(self.modulo.ImageGrab, "grab"), \
             patch.object(self.modulo.time, "sleep"), \
             patch.object(self.modulo, "_executar_ocr", side_effect=[
                 [{"text": "Nome da Empresa"}], [{"text": "Lojas Hoje"}]
             ]) as ocr:
            self.assertTrue(self.modulo._aguardar_empresas_acesso(123, 10))
        self.assertEqual(ocr.call_count, 2)

    def test_login_sem_empresa_interrompe_por_timeout(self):
        with patch.object(self.modulo.time, "monotonic", side_effect=[0, 11]):
            with self.assertRaisesRegex(self.modulo.GestorError, "nenhum novo clique"):
                self.modulo._aguardar_empresas_acesso(123, 10)

    def test_controle_de_acesso_transitorio_reabre_comercial_uma_vez(self):
        config = SimpleNamespace(
            timeout_janela=20,
            usuario="JUNIORM",
            senha="SENHA",
            intervalo_tecla=0.1,
        )
        acesso_primeiro = {"hwnd": 101}
        acesso_segundo = {"hwnd": 102}
        local = {"hwnd": 201}
        with (
            patch.object(self.modulo, "_estado_janelas", return_value={}),
            patch.object(self.modulo, "_clicar_modulo_comercial") as clicar_comercial,
            patch.object(
                self.modulo,
                "aguardar_janela",
                side_effect=[acesso_primeiro, acesso_segundo, local],
            ),
            patch.object(self.modulo, "preparar_janela_gestor"),
            patch.object(
                self.modulo,
                "_aguardar_formulario_acesso_pronto",
                side_effect=[
                    self.modulo.GestorError("Controle de acesso", "janela fechou"),
                    None,
                ],
            ),
            patch.object(self.modulo, "_preparar_teclado_credenciais"),
            patch.object(self.modulo, "_capturar_campo_canvas_acesso", side_effect=["antes", "depois", "antes", "depois"]),
            patch.object(self.modulo, "_preencher_campo_canvas_acesso"),
            patch.object(self.modulo, "_validar_mudanca_visual_campo"),
            patch.object(self.modulo, "_clicar_canvas_acesso"),
            patch.object(self.modulo, "_aguardar_empresas_acesso", return_value=False),
            patch.object(self.modulo, "_aguardar_fechar"),
            patch.object(self.modulo.time, "sleep"),
        ):
            resultado = self.modulo.autenticar_comercial({"hwnd": 1}, config, lambda _: None)

        self.assertEqual(resultado, local)
        self.assertEqual(clicar_comercial.call_count, 3)

    def test_diagnostico_modal_usa_captura_sem_ativar_janela(self):
        with tempfile.TemporaryDirectory() as pasta, \
             patch.object(self.modulo, "_janela_existe", return_value=True), \
             patch.object(self.modulo, "_enumerar_janelas", return_value=[]), \
             patch.object(self.modulo, "_ativar", side_effect=RuntimeError("modal")), \
             patch.object(self.modulo.ImageGrab, "grab") as capturar:
            config = SimpleNamespace(pasta_erros=Path(pasta))
            caminho = self.modulo.salvar_diagnostico_gestor(config, "login", lambda _: None, hwnd=123)
            self.assertIsNotNone(caminho)
            capturar.assert_called_once_with(all_screens=True)

    def test_total_precisa_conferir(self):
        soma = self.modulo.validar_total_relatorio(
            self.valores, Decimal("29866.53")
        )
        self.assertEqual(soma, Decimal("29866.53"))
        with self.assertRaises(self.modulo.GestorError):
            self.modulo.validar_total_relatorio(
                self.valores, Decimal("29866.54")
            )

    def test_parser_separa_meta_de_valor_realizado(self):
        linha = "05 LOJAS HOJE -TUC 105.500,00 2.538,00 2,406"
        self.assertEqual(
            self.modulo._valores_monetarios_linha(linha),
            [Decimal("105500.00"), Decimal("2538.00")],
        )

    def test_ocr_seleciona_apenas_comercial_em_linha_com_varios_modulos(self):
        linha = {
            "texto": "Financeiro Comercial Logoff",
            "palavras": [
                {"text": "Financeiro", "x": 100, "y": 20, "width": 80, "height": 20},
                {"text": "Comercial", "x": 500, "y": 20, "width": 75, "height": 20},
                {"text": "Logoff", "x": 900, "y": 20, "width": 55, "height": 20},
            ],
        }

        palavras = self.modulo._palavras_do_alvo_ocr(linha, "Comercial")

        self.assertEqual([palavra["text"] for palavra in palavras], ["Comercial"])

    def test_ocr_seleciona_frase_contigua_sem_incluir_opcoes_vizinhas(self):
        linha = {
            "texto": "Gestor Lojas Hoje Logoff",
            "palavras": [
                {"text": "Gestor", "x": 10},
                {"text": "Lojas", "x": 100},
                {"text": "Hoje", "x": 150},
                {"text": "Logoff", "x": 500},
            ],
        }

        palavras = self.modulo._palavras_do_alvo_ocr(linha, "Lojas Hoje")

        self.assertEqual([palavra["text"] for palavra in palavras], ["Lojas", "Hoje"])

    def test_formulario_acesso_exige_todos_os_sinais_visuais(self):
        completo = [
            {"text": "Usuário:"},
            {"text": "Senha:"},
            {"text": "Acessar"},
        ]
        incompleto = [{"text": "Usuário:"}, {"text": "Senha:"}]

        self.assertTrue(self.modulo._formulario_acesso_visivel(completo))
        self.assertFalse(self.modulo._formulario_acesso_visivel(incompleto))

    def test_data_campo_aceita_mascara_e_recusa_outro_dia(self):
        data = datetime(2026, 9, 21)

        self.assertTrue(self.modulo._data_campo_corresponde("21/09/2026", data))
        self.assertTrue(self.modulo._data_campo_corresponde("21092026", data))
        self.assertTrue(self.modulo._data_campo_corresponde("211092026", data))
        self.assertTrue(self.modulo._data_campo_corresponde("2110912026", data))
        self.assertFalse(self.modulo._data_campo_corresponde("20/09/2026", data))

    def test_comparacao_visual_do_campo_tolera_ruido_e_detecta_data_alterada(self):
        antes = Image.new("RGB", (100, 30), "white")
        igual = antes.copy()
        igual.putpixel((0, 0), (253, 253, 253))
        alterado = antes.copy()
        ImageDraw.Draw(alterado).rectangle((20, 8, 80, 22), fill="black")

        self.assertTrue(self.modulo._campo_data_permaneceu_igual(antes, igual))
        self.assertFalse(self.modulo._campo_data_permaneceu_igual(antes, alterado))

    def test_leitura_data_inclui_contraste_alternativo_sem_trocar_digitos(self):
        from PIL import Image

        with (
            patch.object(self.modulo, "_ativar"),
            patch.object(self.modulo.time, "sleep"),
            patch.object(self.modulo, "_retangulo_visivel", return_value=(0, 0, 1322, 806)),
            patch.object(self.modulo.ImageGrab, "grab", return_value=Image.new("RGB", (94, 28), "white")),
            patch.object(self.modulo, "_executar_ocr", side_effect=[
                [{"text": "23032026"}], [{"text": "23032026"}], [],
                [{"text": "25/092026"}],
            ]) as ocr,
        ):
            lido = self.modulo._ler_data_visual_campo(1, 0.228, 0.854)

        self.assertEqual(ocr.call_count, 4)
        self.assertTrue(self.modulo._data_campo_corresponde(lido, datetime(2026, 9, 25)))
        self.assertFalse(self.modulo._data_campo_corresponde(lido, datetime(2026, 9, 26)))
        self.assertFalse(self.modulo._data_campo_corresponde("23032026 | 23032026", datetime(2026, 9, 25)))

    def test_data_nao_confirmada_interrompe_sem_aceitar_valor_esperado(self):
        with (
            patch.dict(sys.modules, {"pyautogui": unittest.mock.MagicMock()}),
            patch.object(self.modulo.time, "sleep"),
            patch.object(self.modulo, "_preencher_campo_canvas_acesso") as preencher,
            patch.object(self.modulo, "_ler_data_visual_campo", return_value="23/03/2026"),
        ):
            with self.assertRaisesRegex(self.modulo.GestorError, "nao foi confirmada"):
                self.modulo._preencher_validar_data_canvas(
                    1, 0.228, 0.854, datetime(2026, 9, 25), 0.12, "Data inicial"
                )
        self.assertEqual(preencher.call_count, 2)

    def test_ocr_tolera_um_caractere_perdido_em_rotulo_longo(self):
        linha = {
            "texto": "Relatrio",
            "palavras": [{"text": "Relatrio", "x": 10}],
        }

        palavras = self.modulo._palavras_do_alvo_ocr(linha, "Relatorio")

        self.assertEqual([palavra["text"] for palavra in palavras], ["Relatrio"])

    def test_coordenada_relativa_usa_retangulo_real_da_janela(self):
        with patch.object(
            self.modulo, "_retangulo_visivel", return_value=(300, 100, 1300, 900)
        ):
            self.assertEqual(self.modulo._ponto_relativo(10, 0.25, 0.75), (550, 700))

    def test_retangulo_precisa_ficar_inteiro_no_monitor(self):
        monitor = (0, 0, 1920, 1080)

        self.assertTrue(self.modulo._retangulo_contido((20, 20, 1900, 1060), monitor))
        self.assertFalse(
            self.modulo._retangulo_contido((1000, 20, 2100, 1060), monitor)
        )

    def test_aguardar_janela_aceita_rail_reutilizada_quando_titulo_muda(self):
        janela = {
            "hwnd": 55,
            "titulo": "Impressao de Relatorio (Remoto)",
            "classe": "RAIL_WINDOW",
            "retangulo": (100, 100, 600, 400),
        }
        with (
            patch.object(self.modulo, "_enumerar_janelas", return_value=[janela]),
            patch.object(self.modulo, "_janela_ativa", return_value=55),
        ):
            encontrada = self.modulo.aguardar_janela(
                ("impressao", "relatorio"),
                1,
                novos_desde={55: "gestor - metas de vendas"},
            )

        self.assertEqual(encontrada, janela)

    def test_relatorio_metas_preenche_periodo_inferior(self):
        config = SimpleNamespace(intervalo_tecla=0.1, timeout_janela=10, timeout_relatorio=20)
        metas = {"hwnd": 100}
        janelas = iter(({"hwnd": 200}, {"hwnd": 300}))
        agora = datetime(2026, 9, 21)

        with (
            patch.object(self.modulo, "preparar_janela_gestor"),
            patch.object(self.modulo, "_ativar"),
            patch.object(self.modulo, "_clicar_retangulo_real") as clicar_real,
            patch.object(self.modulo, "_clicar_texto_ocr_janela") as clicar_texto,
            patch.object(
                self.modulo,
                "_localizar_texto_ocr_janela",
                return_value={"x": 560, "y": 900, "altura": 11},
            ),
            patch.object(
                self.modulo, "_retangulo_visivel", return_value=(0, 0, 1000, 1000)
            ),
            patch.object(
                self.modulo,
                "_preencher_validar_data_canvas",
                return_value="21/09/2026",
            ) as preencher,
            patch.object(
                self.modulo, "_ler_data_visual_campo", return_value="21/09/2026"
            ),
            patch.object(
                self.modulo,
                "_capturar_campo_data_canvas",
                return_value=Image.new("RGB", (100, 30), "white"),
            ),
            patch.object(self.modulo, "_estado_janelas", return_value={}),
            patch.object(self.modulo, "_clicar"),
            patch.object(self.modulo, "aguardar_janela", side_effect=lambda *a, **k: next(janelas)),
            patch.object(self.modulo.time, "sleep"),
        ):
            self.modulo.gerar_relatorio_metas(
                metas, config, lambda mensagem: None, agora
            )

        self.assertEqual(
            preencher.call_args_list,
            [
                call(100, 0.228, 0.845, agora, 0.1, "Data inicial"),
                call(100, 0.312, 0.845, agora, 0.1, "Data final"),
            ],
        )
        self.assertEqual(
            clicar_real.call_args_list,
            [
                call(200, 0.07, 0.31),
                call(200, 0.405, 0.915),
            ],
        )
        self.assertEqual(
            clicar_texto.call_args_list,
            [
                call(100, "Relatorio", ANY, faixa_y=(0.75, 0.99)),
                call(100, "Iniciar", ANY, faixa_y=(0.75, 0.99)),
            ],
        )

    def test_excel_altera_somente_c19_c26(self):
        with tempfile.TemporaryDirectory() as pasta:
            pasta = Path(pasta)
            arquivo = pasta / "metas.xlsx"
            criar_planilha(arquivo)
            antes = load_workbook(arquivo, data_only=False)
            formulas = [antes[f"base"][f"D{linha}"].value for linha in range(19, 27)]
            antes.close()

            gravados, backup = self.modulo.atualizar_excel_metas(
                arquivo, self.valores, pasta / "backups"
            )

            self.assertTrue(backup.exists())
            self.assertEqual(set(gravados), set(self.modulo.LOJAS))
            depois = load_workbook(arquivo, data_only=False)
            try:
                ws = depois["base"]
                for loja, linha in self.modulo.LINHAS_EXCEL.items():
                    self.assertEqual(
                        Decimal(str(ws[f"C{linha}"].value)).quantize(Decimal("0.01")),
                        self.valores[loja],
                    )
                    self.assertEqual(ws[f"C{linha}"].fill.fgColor.rgb, "0000FF00")
                self.assertEqual(
                    [ws[f"D{linha}"].value for linha in range(19, 27)], formulas
                )
                self.assertEqual(ws["Z50"].value, "nao alterar")
                self.assertEqual(depois["outra"]["A1"].value, "preservado")
            finally:
                depois.close()

    def test_excel_recusa_rotulo_de_loja_incorreto(self):
        with tempfile.TemporaryDirectory() as pasta:
            arquivo = Path(pasta) / "metas.xlsx"
            criar_planilha(arquivo, rotulo_errado=True)
            with self.assertRaises(self.modulo.GestorError):
                self.modulo.atualizar_excel_metas(
                    arquivo, self.valores, Path(pasta) / "backups"
                )

    def test_excel_aceita_rotulos_por_referencia_simples(self):
        with tempfile.TemporaryDirectory() as pasta:
            arquivo = Path(pasta) / "metas.xlsx"
            criar_planilha(arquivo, rotulos_por_formula=True)

            gravados, _backup = self.modulo.atualizar_excel_metas(
                arquivo, self.valores, Path(pasta) / "backups"
            )

            self.assertEqual(set(gravados), set(self.modulo.LOJAS))


if __name__ == "__main__":
    unittest.main()
